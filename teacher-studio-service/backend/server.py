from __future__ import annotations

import json
import base64
import time
import threading
import io
import html
import hashlib
import os
import re
import shutil
import ssl
import struct
import subprocess
import tempfile
import uuid
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from email.parser import BytesParser
from email.policy import default
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import quote, unquote, urlparse, urljoin
from urllib.request import Request, urlopen

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
import biomed_curriculum
import biomed_models
import biomed_slides
import courseware
import heart_anatomy
import image_model
import lesson_design
import local_reconstruction
import lesson_ai_adapter
import partfield_adapter
import triposr_local
import teaching_benchmarks
from express_model import obj_archive_to_glb


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "generated"
OUTPUT_DIR.mkdir(exist_ok=True)
ASSET_DIR = OUTPUT_DIR / "assets"
ASSET_DIR.mkdir(exist_ok=True)


def env_value(name, fallback=""):
    value = os.environ.get(name)
    if value is not None:
        return value.strip()
    env_files = (ROOT / ".env", ROOT.parent / ".env.local", ROOT.parent / ".env")
    for env_file in env_files:
        if not env_file.exists():
            continue
        for line in env_file.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, file_value = line.split("=", 1)
            if key.strip() == name:
                return file_value.strip().strip('"').strip("'")
    return fallback


def now_iso():
    return datetime.now().isoformat(timespec="seconds")


def generate_warmup_questions(payload):
    """Build five local, editable questions from the current lesson and model manifest."""
    lesson = payload.get("lessonPlan") or {}
    model = payload.get("model") or {}
    model_url = str(model.get("glbUrl") or "").strip() if isinstance(model, dict) else ""
    parts = (model.get("parts") or model.get("partDetails") or []) if isinstance(model, dict) else []
    normalized = []
    for part in parts:
        if isinstance(part, str):
            normalized.append({"key": part, "name": part})
        elif isinstance(part, dict):
            key = part.get("partKey") or part.get("key") or part.get("name")
            name = part.get("partLabel") or part.get("name") or key
            if key and name:
                normalized.append({"key": key, "name": name})
    normalized = normalized[:12]
    has_model = bool(model_url and normalized)
    lesson_id = str(lesson.get("id") or uuid.uuid4().hex)
    title = str(lesson.get("title") or "本课内容")
    subject = str((lesson.get("request") or {}).get("subject") or "生物")
    objectives = lesson.get("objectives") or lesson.get("learningOutcomes") or []
    objective = str(objectives[0]) if objectives else f"解释{title}中的核心结构与机制"

    if has_model:
        names = [item["name"] for item in normalized]
        keys = [item["key"] for item in normalized]
        first = normalized[0]
        second = normalized[1] if len(normalized) > 1 else normalized[0]

        def model_options(correct):
            values = [correct]
            for candidate in names + ["其他未标注结构", "模型外部背景", "无法仅凭颜色判断"]:
                if candidate not in values:
                    values.append(candidate)
                if len(values) == 4:
                    break
            return values

        questions = [
            {"question": f"在“{title}”模型中，观察任务首先要求定位哪个部件？", "options": model_options(first["name"]), "correctIndex": 0, "partKeys": [first["key"]], "difficulty": "基础", "explanation": f"应先依据部件树和空间位置定位{first['name']}，再观察其连接关系。"},
            {"question": f"观察{second['name']}时，判断其位置或连接关系最可靠的依据是什么？", "options": ["部件名称、空间位置与相邻结构", "模型背景颜色", "屏幕亮度", "文件大小"], "correctIndex": 0, "partKeys": [second["key"]], "difficulty": "基础", "explanation": "结构关系应由名称、位置和相邻连接共同判断。"},
            {"question": "旋转、隐藏或隔离模型部件后，哪项观察记录最有效？", "options": ["记录视角变化前后的空间关系", "只记录部件颜色", "只记录缩放倍数", "忽略相邻结构"], "correctIndex": 0, "partKeys": keys[:2], "difficulty": "基础", "explanation": "改变视角是为了核对结构的三维空间关系，而不是比较显示颜色。"},
            {"question": "观察模型时，怎样从结构推断其功能？", "options": ["结合位置、形态、连接关系与机制解释", "只根据外观颜色判断", "只根据部件大小判断", "不需要证据直接猜测"], "correctIndex": 0, "partKeys": keys[:2], "difficulty": "进阶", "explanation": "结构功能推断需要同时使用形态、位置、连接关系和课程机制。"},
            {"question": f"针对“{title}”，形成课堂结论时最完整的做法是什么？", "options": ["说明结构、机制、证据及适用边界", "只写一个部件名称", "只复述模型颜色", "跳过证据直接下结论"], "correctIndex": 0, "partKeys": keys[:3], "difficulty": "进阶", "explanation": "完整结论需要结构事实、机制解释和证据边界相互支持。"},
        ]
    else:
        questions = [
            {"question": f"阅读“{title}”结构图时，第一步应完成什么任务？", "options": ["确认图例、方向和关键结构名称", "只看配色是否美观", "跳过标注直接记答案", "只统计结构数量"], "correctIndex": 0, "partKeys": [], "difficulty": "基础", "explanation": "结构图观察应先确认图例、方向和标注，再分析结构关系。"},
            {"question": f"分析“{title}”中的结构位置或连接关系，哪类证据最可靠？", "options": ["教材标注、相邻关系与流程方向", "页面背景颜色", "文字字号大小", "图片装饰效果"], "correctIndex": 0, "partKeys": [], "difficulty": "基础", "explanation": "位置与连接关系应由规范标注、相邻结构和过程方向共同证明。"},
            {"question": f"解释“{title}”的机制过程时，合理的分析顺序是什么？", "options": ["条件或输入→关键过程→结果或输出", "先写结论再删除证据", "只罗列名词不说明关系", "按页面颜色排列"], "correctIndex": 0, "partKeys": [], "difficulty": "基础", "explanation": "机制解释需要呈现条件、关键过程与结果之间的因果链。"},
            {"question": f"要完成学习目标“{objective}”，哪种回答最充分？", "options": ["用结构证据解释功能或机制", "只背诵一个术语", "只描述图片颜色", "省略理由只给结论"], "correctIndex": 0, "partKeys": [], "difficulty": "进阶", "explanation": "有效回答应把结构证据与功能或机制联系起来。"},
            {"question": f"将“{title}”应用到新案例时，最合适的判断方法是什么？", "options": ["提取案例证据并逐项对应课堂机制", "凭第一印象直接选择", "只寻找熟悉词语", "忽略不支持结论的信息"], "correctIndex": 0, "partKeys": [], "difficulty": "进阶", "explanation": "迁移应用要求把案例证据与课堂机制逐项对应，并检查反例。"},
        ]
    _spread_warmup_answers(questions)
    alignment = lesson.get("alignmentMatrix") or []
    result = []
    for index, item in enumerate(questions):
        alignment_item = alignment[index % len(alignment)] if alignment and isinstance(alignment[index % len(alignment)], dict) else {}
        objective_id = alignment_item.get("objectiveId")
        activity_ids = alignment_item.get("activityIds") or []
        result.append({
            "id": f"warmup-{lesson_id}-{index + 1}",
            "modelUrl": model_url if has_model else "",
            "subject": subject,
            "category": "model-warmup",
            "question": item["question"],
            "options": item["options"],
            "correctIndex": item["correctIndex"],
            "explanation": item["explanation"],
            "optionType": 4,
            "modelPartKeys": item["partKeys"],
            "lessonId": lesson_id,
            "difficulty": item["difficulty"],
            "enabled": True,
            "source": "generated",
            "objectiveIds": [objective_id] if objective_id else [],
            "activityId": activity_ids[0] if activity_ids else "",
        })
    return result


WARMUP_TYPES = ("部件识别", "位置或连接关系", "模型观察判断", "结构与功能", "机制迁移")
WARMUP_CORRECT_POSITIONS = (0, 1, 2, 3, 1)


def _spread_warmup_answers(items, indexes=None):
    indexes = indexes or WARMUP_CORRECT_POSITIONS
    for index, item in enumerate(items):
        options = item.get("options")
        if not isinstance(options, list) or len(options) < 4:
            continue
        target = indexes[index % len(indexes)]
        current = item.get("correctIndex", 0)
        if isinstance(current, int) and 0 <= current < len(options) and current != target:
            options[current], options[target] = options[target], options[current]
            item["correctIndex"] = target
    return items


def _warmup_model_context(model):
    model = model if isinstance(model, dict) else {}
    model_url = str(model.get("glbUrl") or model.get("modelUrl") or "").strip()
    parts = model.get("parts") or model.get("partDetails") or []
    normalized = []
    for part in parts:
        if isinstance(part, str):
            normalized.append({"key": part, "name": part})
        elif isinstance(part, dict):
            key = part.get("partKey") or part.get("key") or part.get("name")
            name = part.get("partLabel") or part.get("name") or key
            if key and name:
                normalized.append({"key": str(key), "name": str(name)})
    return model_url, normalized[:16]


def _warmup_question_schema():
    return {
        "type": "object",
        "properties": {
            "questions": {
                "type": "array",
                "minItems": 5,
                "maxItems": 5,
                "items": {
                    "type": "object",
                    "properties": {
                        "type": {"type": "string"},
                        "question": {"type": "string"},
                        "options": {"type": "array", "minItems": 4, "maxItems": 4, "items": {"type": "string"}},
                        "correctIndex": {"type": "integer"},
                        "explanation": {"type": "string"},
                        "difficulty": {"type": "string"},
                        "partKeys": {"type": "array", "items": {"type": "string"}},
                    },
                    "required": ["type", "question", "options", "correctIndex", "explanation", "difficulty", "partKeys"],
                    "additionalProperties": False,
                },
            },
        },
        "required": ["questions"],
        "additionalProperties": False,
    }


def _validate_warmup_ai(value, valid_part_keys, has_model, expected_types=None):
    questions = value.get("questions") if isinstance(value, dict) else None
    if not isinstance(questions, list) or len(questions) != 5:
        return "必须返回恰好 5 道题"
    expected_types = expected_types or WARMUP_TYPES
    for index, item in enumerate(questions):
        if not isinstance(item, dict):
            return f"第 {index + 1} 题不是对象"
        if item.get("type") != expected_types[index]:
            return f"第 {index + 1} 题类型必须是“{expected_types[index]}”"
        question = str(item.get("question") or "").strip()
        explanation = str(item.get("explanation") or "").strip()
        options = item.get("options")
        if not question or not explanation:
            return f"第 {index + 1} 题题干和解析不能为空"
        if not isinstance(options, list) or len(options) != 4 or any(not str(option).strip() for option in options):
            return f"第 {index + 1} 题必须有 4 个有效选项"
        if len({str(option).strip() for option in options}) != 4:
            return f"第 {index + 1} 题不能有重复选项"
        if not isinstance(item.get("correctIndex"), int) or not 0 <= item["correctIndex"] < 4:
            return f"第 {index + 1} 题正确答案无效"
        if item.get("difficulty") not in ("基础", "进阶"):
            return f"第 {index + 1} 题难度必须是基础或进阶"
        part_keys = item.get("partKeys")
        if not isinstance(part_keys, list) or any(key not in valid_part_keys for key in part_keys):
            return f"第 {index + 1} 题引用了不存在的模型部件"
        if not has_model and part_keys:
            return f"无模型课程的第 {index + 1} 题不能引用模型部件"
    return None


def _warmup_with_metadata(items, lesson, model_url, generation):
    alignment = lesson.get("alignmentMatrix") or []
    result = []
    for index, item in enumerate(items[:5]):
        part_keys = item.get("partKeys") or item.get("modelPartKeys") or []
        alignment_item = alignment[index % len(alignment)] if alignment and isinstance(alignment[index % len(alignment)], dict) else {}
        objective_id = alignment_item.get("objectiveId")
        activity_ids = alignment_item.get("activityIds") or []
        result.append({
            "id": item.get("id") or f"warmup-{lesson.get('id') or uuid.uuid4().hex}-{index + 1}",
            "modelUrl": model_url if part_keys else "",
            "subject": str((lesson.get("request") or {}).get("subject") or "生物"),
            "category": "model-warmup",
            "question": item["question"],
            "options": item["options"],
            "correctIndex": item["correctIndex"],
            "explanation": item["explanation"],
            "optionType": 4,
            "modelPartKeys": part_keys,
            "lessonId": str(lesson.get("id") or ""),
            "difficulty": item["difficulty"],
            "enabled": item.get("enabled", True),
            "source": item.get("source", "generated"),
            "objectiveIds": item.get("objectiveIds") or ([objective_id] if objective_id else []),
            "activityId": item.get("activityId") or (activity_ids[0] if activity_ids else ""),
        })
    return result


def generate_warmup_with_ai(payload, current_questions=None, regenerate_index=None):
    lesson = payload.get("lessonPlan") or {}
    model_url, normalized_parts = _warmup_model_context(payload.get("model"))
    has_model = bool(model_url and normalized_parts)
    valid_part_keys = {item["key"] for item in normalized_parts}
    template = generate_warmup_questions(payload)
    teacher_questions = current_questions if isinstance(current_questions, list) else []
    expected_types = list(WARMUP_TYPES)
    if regenerate_index is not None:
        regenerate_index = max(0, min(4, int(regenerate_index)))
        expected_types = [WARMUP_TYPES[regenerate_index]]

    model_summary = [{"key": item["key"], "name": item["name"]} for item in normalized_parts]
    questions_context = [
        {"question": item.get("question"), "type": WARMUP_TYPES[index] if index < 5 else ""}
        for index, item in enumerate(teacher_questions)
        if index != regenerate_index
    ]
    prompt = {
        "task": "生成练习题，不是正式考试题",
        "course": {
            "title": lesson.get("title"),
            "subject": (lesson.get("request") or {}).get("subject"),
            "grade": (lesson.get("request") or {}).get("grade"),
            "abilityLevel": (lesson.get("request") or {}).get("abilityLevel"),
            "duration": (lesson.get("request") or {}).get("duration"),
        },
        "objectives": lesson.get("objectives") or lesson.get("learningOutcomes") or [],
        "activities": lesson.get("activities") or [],
        "alignmentMatrix": lesson.get("alignmentMatrix") or [],
        "sourceMaterial": str((lesson.get("request") or {}).get("materialText") or "")[:6000],
        "modelParts": model_summary if has_model else [],
        "requiredTypes": expected_types,
        "avoidRepeating": questions_context,
        "rules": [
            "每题必须有 4 个互不重复的选项和 1 个正确答案",
            "解析必须说明为什么正确，并服务于课堂练习",
            "没有模型时 partKeys 必须为空，不能编造部件",
            "只使用提供的课程资料、教案和模型部件，不补写未经依据的医学事实",
        ],
    }

    def validate(value):
        return _validate_warmup_ai(value, valid_part_keys, has_model, expected_types)

    result, error = lesson_ai_adapter.generate_json(
        "你是严谨的生物医学教师。生成适合课堂开始阶段的可编辑选择题。只返回符合要求的 JSON，不输出解释文字。",
        json.dumps(prompt, ensure_ascii=False),
        validator=validate,
        timeout=180,
        format_schema=_warmup_question_schema(),
    )
    if result is None:
        fallback = template[regenerate_index:regenerate_index + 1] if regenerate_index is not None else template
        fallback_generation = {"engine": "builtin-template", "model": lesson_ai_adapter.DEFAULT_MODEL, "message": f"远程 AI 生成失败，模板回退：{error}"}
        return fallback, fallback_generation, [fallback_generation["message"]]

    ai_items = []
    for item in result["questions"]:
        ai_items.append({
            "question": item["question"].strip(),
            "options": [str(option).strip() for option in item["options"]],
            "correctIndex": item["correctIndex"],
            "explanation": item["explanation"].strip(),
            "partKeys": item.get("partKeys") or [],
            "difficulty": item["difficulty"],
        })
    if regenerate_index is not None:
        _spread_warmup_answers(ai_items, WARMUP_CORRECT_POSITIONS[regenerate_index:regenerate_index + 1])
        generated = template.copy()
        generated[regenerate_index] = ai_items[0]
        return generated, {"engine": "responses", "model": lesson_ai_adapter.DEFAULT_MODEL, "message": f"已使用 {lesson_ai_adapter.DEFAULT_MODEL} 重新生成第 {regenerate_index + 1} 题"}, []
    _spread_warmup_answers(ai_items)
    merged = []
    for index, item in enumerate(ai_items):
        if index < len(teacher_questions) and teacher_questions[index].get("source") == "teacher-edited":
            merged.append(teacher_questions[index])
        else:
            merged.append(item)
    warnings = ["已保留教师编辑题目，未自动覆盖"] if any(item.get("source") == "teacher-edited" for item in teacher_questions) else []
    return merged, {"engine": "responses", "model": lesson_ai_adapter.DEFAULT_MODEL, "message": f"已使用 {lesson_ai_adapter.DEFAULT_MODEL} 生成 5 道练习题"}, warnings


def parse_multipart(body, content_type, allow_multiple=False):
    message = BytesParser(policy=default).parsebytes(
        f"Content-Type: {content_type}\r\nMIME-Version: 1.0\r\n\r\n".encode("utf-8") + body
    )
    if not message.is_multipart():
        raise ValueError("3D 请求必须使用 multipart/form-data")
    fields = {}
    images = []
    for part in message.iter_parts():
        disposition = part.get("Content-Disposition", "")
        name = part.get_param("name", header="content-disposition")
        if not name:
            continue
        value = part.get_payload(decode=True) or b""
        filename = part.get_filename()
        if filename and (name == "image" or name == "images" or allow_multiple):
            images.append({"field": name, "filename": Path(filename).name or "image.png", "content": value, "content_type": part.get_content_type()})
        elif not filename:
            fields[name] = value.decode(part.get_content_charset() or "utf-8", errors="replace")
    if not images or not any(image["content"] for image in images):
        raise ValueError("请上传一张 PNG 或 JPG 图片")
    if allow_multiple:
        return fields, images
    return fields, images[0]


def build_multipart(fields, image):
    boundary = f"----TeacherStudio{uuid.uuid4().hex}"
    chunks = []
    for name, value in fields.items():
        chunks.extend([
            f"--{boundary}\r\n".encode(),
            f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode(),
            str(value).encode(),
            b"\r\n",
        ])
    chunks.extend([
        f"--{boundary}\r\n".encode(),
        f'Content-Disposition: form-data; name="image"; filename="{image["filename"]}"\r\n'.encode(),
        f'Content-Type: {image["content_type"] or "application/octet-stream"}\r\n\r\n'.encode(),
        image["content"],
        b"\r\n",
        f"--{boundary}--\r\n".encode(),
    ])
    return b"".join(chunks), f"multipart/form-data; boundary={boundary}"


def build_multipart_many(fields, images):
    boundary = f"----TeacherStudio{uuid.uuid4().hex}"
    chunks = []
    for name, value in fields.items():
        chunks.extend([
            f"--{boundary}\r\n".encode(),
            f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode(),
            str(value).encode(), b"\r\n",
        ])
    for image in images:
        chunks.extend([
            f"--{boundary}\r\n".encode(),
            f'Content-Disposition: form-data; name="{image["field"]}"; filename="{image["filename"]}"\r\n'.encode(),
            f'Content-Type: {image["content_type"] or "application/octet-stream"}\r\n\r\n'.encode(),
            image["content"], b"\r\n",
        ])
    chunks.append(f"--{boundary}--\r\n".encode())
    return b"".join(chunks), f"multipart/form-data; boundary={boundary}"


MULTIVIEW_KEYS = ("front", "left", "right", "back")
MULTIVIEW_LABELS = {"front": "正面", "left": "左侧", "right": "右侧", "back": "背面"}


def hunyuan3d_config():
    return (
        env_value("HUNYUAN3D_BASE_URL", "https://tokenhub.tencentmaas.com").rstrip("/"),
        env_value("HUNYUAN3D_API_KEY"),
        env_value("HUNYUAN3D_MODEL", "hy-3d-3.1"),
    )


def hunyuan3d_model_config(requested_model=""):
    base_url, default_key, default_model = hunyuan3d_config()
    model = str(requested_model or default_model).strip()
    if model == "hy-3d-express":
        return base_url, env_value("HUNYUAN3D_EXPRESS_API_KEY", default_key), model
    if model == "hy-3d-3.1":
        return base_url, default_key, model
    raise ValueError("混元 3D 模型版本无效")


def _hunyuan_json_request(path, payload, api_key, timeout=60):
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    request = Request(
        path,
        data=body,
        headers={"Content-Type": "application/json", "Authorization": f"Bearer {api_key}", "Accept": "application/json"},
        method="POST",
    )
    try:
        with urlopen(request, timeout=timeout) as response:
            raw = response.read()
    except HTTPError as error:
        detail = error.read().decode("utf-8", errors="replace")[-1200:]
        raise RuntimeError(f"混元 3D 接口失败（HTTP {error.code}）：{detail.replace(api_key, '[redacted]')}") from error
    except URLError as error:
        raise RuntimeError(f"混元 3D 服务不可达：{error.reason}") from error
    try:
        result = json.loads(raw.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as error:
        raise RuntimeError("混元 3D 返回了无法解析的响应") from error
    if not isinstance(result, dict):
        raise RuntimeError("混元 3D 返回格式无效")
    return result


def _hunyuan_download_glb(url, archive=False):
    parsed = urlparse(url)
    if parsed.scheme != 'https' or not parsed.hostname or not parsed.hostname.endswith(('.tencentcos.cn', '.myqcloud.com')) or parsed.username or parsed.password:
        raise RuntimeError('混元返回了非允许的模型下载地址，已阻止下载')
    last_error = None
    for attempt in range(1, 4):
        try:
            context = ssl.create_default_context()
            # Some Tencent COS edge nodes intermittently close TLS 1.3 handshakes
            # from the bundled Windows/OpenSSL runtime; keep certificate checks
            # enabled and use TLS 1.2 for the fallback connection.
            if attempt > 1:
                context.maximum_version = ssl.TLSVersion.TLSv1_2
            request = Request(url, headers={
                "Accept": "application/octet-stream",
                "Connection": "close",
            })
            # COS download URLs occasionally close an idle TLS connection before
            # the body arrives. A fresh verified connection makes this recoverable.
            with urlopen(request, timeout=120, context=context) as response:
                result = response.read()
            if archive:
                return obj_archive_to_glb(result)
            if len(result) < 12 or not result.startswith(b"glTF") or struct.unpack_from("<II", result, 4) != (2, len(result)):
                raise RuntimeError("混元 3D 返回的文件不是完整有效的 GLB")
            return result
        except HTTPError as error:
            last_error = error
            if error.code < 500 or attempt == 3:
                if error.code < 500:
                    detail = error.read().decode("utf-8", errors="replace")[-600:]
                    raise RuntimeError(f"混元 3D 模型文件下载失败（HTTP {error.code}）：{detail}") from error
        except (URLError, TimeoutError, ssl.SSLError) as error:
            last_error = error
        time.sleep(attempt * 1.5)
    # Signed URLs are passed as data, never interpolated into shell commands.
    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(prefix="hunyuan-", suffix=".glb", delete=False) as temporary:
            temporary_path = temporary.name
        native_environment = os.environ.copy()
        native_environment["HUNYUAN_DOWNLOAD_URL"] = url
        native_environment["HUNYUAN_DOWNLOAD_OUT"] = temporary_path
        native_environment["HUNYUAN_DOWNLOAD_FORMAT"] = "zip" if archive else "glb"
        node = shutil.which("node.exe") or shutil.which("node")
        if not node:
            raise RuntimeError("未找到 Node.js，无法执行备用模型下载")
        completed = subprocess.run(
            [node, str(Path(__file__).with_name("download_hunyuan.mjs"))],
            env=native_environment,
            capture_output=True,
            text=True,
            timeout=130,
            check=False,
        )
        if completed.returncode == 0:
            result = Path(temporary_path).read_bytes()
        else:
            detail = (completed.stderr or completed.stdout or str(last_error)).strip()[-600:]
            if archive:
                raise RuntimeError("Express OBJ 下载失败，请重试以恢复已完成任务")
            # The provider output is still valid and signed. Let the browser
            # fetch it directly when this machine cannot reach the COS edge.
            return {"remoteUrl": url, "downloadError": detail}
    except (OSError, subprocess.SubprocessError, RuntimeError) as error:
        raise RuntimeError(f"混元 3D 模型文件下载失败：{error}") from error
    finally:
        if temporary_path:
            Path(temporary_path).unlink(missing_ok=True)
    if archive:
        return obj_archive_to_glb(result)
    if len(result) < 12 or not result.startswith(b"glTF") or struct.unpack_from("<II", result, 4) != (2, len(result)):
        raise RuntimeError("混元 3D 返回的文件不是完整有效的 GLB")
    return result


_hunyuan_slots = threading.BoundedSemaphore(3)


def generate_hunyuan3d(image, fields, view_label, multiview_images=None):
    if not _hunyuan_slots.acquire(blocking=False):
        raise RuntimeError('已有三个混元任务正在处理，请等待完成后再提交')
    try:
        return _generate_hunyuan3d(image, fields, view_label, multiview_images=multiview_images)
    finally:
        _hunyuan_slots.release()


def _encode_hunyuan_image(image, label, max_encoded_bytes=6 * 1024 * 1024):
    try:
        with Image.open(io.BytesIO(image["content"])) as source:
            if source.format not in ("JPEG", "PNG", "WEBP") or not all(128 <= size <= 5000 for size in source.size):
                raise ValueError(f"{label}图片须为 PNG/JPG/WebP，单边尺寸 128 至 5000 像素")
            source.verify()
    except (OSError, ValueError) as error:
        raise ValueError(f"{label}图片无效：{error}") from error

    # Four Base64 images share the provider's 10 MB JSON body limit. Re-encode
    # multiview inputs into bounded JPEGs so users do not need to resize them.
    with Image.open(io.BytesIO(image["content"])) as source:
        source = ImageOps.exif_transpose(source)
        if source.mode in ("RGBA", "LA") or "transparency" in source.info:
            canvas = Image.new("RGB", source.size, "white")
            alpha = source.getchannel("A") if "A" in source.getbands() else None
            canvas.paste(source.convert("RGBA"), mask=alpha)
            source = canvas
        else:
            source = source.convert("RGB")
        source.thumbnail((1536, 1536), Image.Resampling.LANCZOS)
        quality = 84
        encoded = ""
        while quality >= 52:
            output = io.BytesIO()
            source.save(output, format="JPEG", quality=quality, optimize=True, progressive=True)
            encoded = base64.b64encode(output.getvalue()).decode("ascii")
            if len(encoded) <= max_encoded_bytes:
                return encoded
            quality -= 8

        # Extremely detailed images may still exceed the budget after quality
        # reduction; shrink once more and encode at a readable fallback quality.
        source.thumbnail((1024, 1024), Image.Resampling.LANCZOS)
        output = io.BytesIO()
        source.save(output, format="JPEG", quality=60, optimize=True, progressive=True)
        encoded = base64.b64encode(output.getvalue()).decode("ascii")
        if len(encoded) > max_encoded_bytes:
            raise ValueError(f"{label}图片压缩后仍超过混元请求限制，请使用更小的图片")
        return encoded


def _validate_hunyuan_image(image, label, max_encoded_bytes=6 * 1024 * 1024):
    return _encode_hunyuan_image(image, label, max_encoded_bytes=max_encoded_bytes)


def _generate_hunyuan3d(image, fields, view_label, multiview_images=None):
    base_url, api_key, model = hunyuan3d_model_config(fields.get("model") or fields.get("modelKey"))
    if not api_key:
        raise RuntimeError("未配置 HUNYUAN3D_API_KEY，无法调用混元 3D")
    image_budget = 1_900_000 if multiview_images else 6 * 1024 * 1024
    encoded = _validate_hunyuan_image(image, view_label, max_encoded_bytes=image_budget)
    payload = {
        "model": model,
        "image_base64": encoded,
        "enable_pbr": True,
        "generate_type": "Normal",
        "face_count": 1500000,
    }
    if multiview_images:
        payload["multi_view_images"] = [
            {
                "view_type": item["field"],
                "view_image_base64": _validate_hunyuan_image(item, MULTIVIEW_LABELS[item["field"]], max_encoded_bytes=image_budget),
            }
            for item in multiview_images
        ]
    # Retry the existing provider task for identical inputs after download failure.
    fingerprint = hashlib.sha256(json.dumps(payload, sort_keys=True).encode("utf-8")).hexdigest()
    task_directory = OUTPUT_DIR / "hunyuan-tasks"
    task_directory.mkdir(exist_ok=True)
    task_path = task_directory / f"{fingerprint}.json"
    stored_task = json.loads(task_path.read_text(encoding="utf-8")) if task_path.exists() else {}
    if stored_task.get("id") and stored_task.get("fingerprint") == fingerprint:
        response = _hunyuan_json_request(f"{base_url}/v1/api/3d/query", {"model": model, "id": str(stored_task["id"])}, api_key)
    else:
        response = _hunyuan_json_request(f"{base_url}/v1/api/3d/submit", payload, api_key)
    task_id = response.get("id")
    if not task_id:
        task_id = stored_task.get("id")
    if not task_id:
        raise RuntimeError("混元 3D 未返回任务编号")
    task_path.write_text(json.dumps({"id": str(task_id), "model": model, "fingerprint": fingerprint}), encoding="utf-8")
    deadline = time.monotonic() + 900
    status = response.get("status", "queued")
    while status in {"queued", "in_progress"}:
        if time.monotonic() >= deadline:
            raise RuntimeError(f"混元 3D {view_label}任务处理超时（任务号：{task_id}）")
        time.sleep(4)
        queried = _hunyuan_json_request(f"{base_url}/v1/api/3d/query", {"model": model, "id": str(task_id)}, api_key)
        status = queried.get("status", "")
        if status == "failed":
            raise RuntimeError(f"混元 3D {view_label}任务失败：{queried.get('message') or queried.get('error') or queried}")
        response = queried
    if status != "completed":
        raise RuntimeError(f"混元 3D {view_label}任务返回未知状态：{status}")
    files = response.get("data") or []
    output = next((item for item in files if str(item.get("type", "")).lower() == "glb" and item.get("url")), None)
    if output is None:
        output = next((item for item in files if str(item.get("type", "")).lower() == "obj" and item.get("url")), None)
    if output is None:
        raise RuntimeError("混元 3D 任务完成，但未返回可用的 GLB 或 OBJ 文件")
    task_path.write_text(json.dumps({"id": str(task_id), "model": model, "fingerprint": fingerprint, "url": output["url"]}), encoding="utf-8")
    if str(output.get("type", "")).lower() == "obj":
        return _hunyuan_download_glb(output["url"], archive=True)
    return _hunyuan_download_glb(output["url"])


def _looks_like_heart(images):
    """Detect the built-in heart reference class without a network classifier."""
    red_scores = []
    blue_scores = []
    for image in images:
        sample = Image.open(io.BytesIO(image["content"])).convert("RGB").resize((64, 64))
        red = blue = 0
        for red_value, green_value, blue_value in sample.getdata():
            red += red_value > max(90, green_value * 1.25, blue_value * 1.15)
            blue += blue_value > max(70, red_value * 1.15, green_value * 1.05)
        red_scores.append(red / 4096)
        blue_scores.append(blue / 4096)
    return sum(red_scores) / len(red_scores) > 0.18 and sum(blue_scores) / len(blue_scores) > 0.015


def parse_multiview(body, content_type):
    fields, images = parse_multipart(body, content_type, allow_multiple=True)
    by_view = {image["field"]: image for image in images if image["field"] in MULTIVIEW_KEYS}
    missing = [key for key in MULTIVIEW_KEYS if not by_view.get(key) or not by_view[key]["content"]]
    if missing:
        raise ValueError("多视角模式需要上传：" + "、".join(missing))
    return fields, [by_view[key] for key in MULTIVIEW_KEYS]


def package_reconstruction(result, source_label, mode, fields, view_payloads=None):
    separation_mode = fields.get("separationMode", fields.get("separation_mode", "geometric"))
    if separation_mode != "anatomical":
        return attach_standard_views(image_model.package_glb(result, source_label, mode))
    organ_type = fields.get("organType", fields.get("organ_type", ""))
    anatomy_profile = fields.get("anatomyProfile", fields.get("anatomy_profile", ""))
    if organ_type != "heart" or anatomy_profile not in ("heart-v1", "heart-v2"):
        raise ValueError("医学语义拆解仅支持 heart / heart-v2（兼容 heart-v1），未执行普通几何回退")
    return attach_standard_views(heart_anatomy.package(result, source_label, mode, view_payloads=view_payloads))


def attach_standard_views(result):
    model_url = result.get('appearanceUrl') or result.get('glbUrl')
    if model_url and model_url.startswith('/api/'):
        try:
            prepared = courseware.prepare_model_views({'model': {'modelUrl': model_url}, 'scenes': [{'model': {}}]})
            result.update({key: value for key, value in prepared.items() if key != 'status'})
            result['viewStatus'] = 'ready'
        except (ValueError, RuntimeError, OSError) as exc:
            result['viewStatus'] = 'failed'
            result['viewError'] = str(exc)
            result['viewImageUrls'] = {}
    return result


def generate_image_model(body, content_type, mode):
    fields, _ = parse_multipart(body, content_type, allow_multiple=True)
    base_url, api_key, model = hunyuan3d_model_config(fields.get("model") or fields.get("modelKey"))
    if not api_key:
        raise RuntimeError("HUNYUAN3D_API_KEY 未配置；建模已锁定混元，不会回退")
    if mode == "single":
        fields, image = parse_multipart(body, content_type)
        result = generate_hunyuan3d(image, fields, "单图")
        views_used = ["front"]
    else:
        fields, images = parse_multiview(body, content_type)
        by_view = {image["field"]: image for image in images}
        result = generate_hunyuan3d(
            by_view["front"],
            fields,
            "正面",
            multiview_images=[by_view[key] for key in ("left", "right", "back")],
        )
        views_used = list(MULTIVIEW_KEYS)
    # Preserve the provider GLB exactly; never assemble anatomy or split regions.
    directory = OUTPUT_DIR / "hunyuan3d" / uuid.uuid4().hex
    directory.mkdir(parents=True)
    path = directory / "model.glb"
    remote_url = result.get("remoteUrl") if isinstance(result, dict) else ""
    if remote_url:
        url = remote_url
        message_suffix = "当前网络无法下载本地副本，已保留混元签名模型地址供浏览器直接加载。"
    else:
        path.write_bytes(result)
        url = "/api/3d/download/" + path.relative_to(OUTPUT_DIR).as_posix()
        message_suffix = "模型文件已保存到本地。"
    packaged = attach_standard_views({
        "status": "generated", "editable": True, "glbUrl": url,
        "appearanceUrl": url, "blendUrl": "", "previewUrl": "",
        "parts": [], "anatomyDisplayAvailable": False,
        "modelKind": "hunyuan3d-appearance", "engine": "hunyuan3d",
        "model": model, "paidApi": True, "sourceApi": base_url,
        "viewsUsed": views_used, "inputViewCount": len(views_used), "qualityTier": "provider-output-unreviewed",
        "message": f"混元 3D 原始模型；已使用{len(views_used)}个视角输入，未添加模板、未自动拆件。{message_suffix}",
        "workflow": ["hunyuan3d", "preserve-original-glb", "four-view-input" if mode != "single" else "single-view-input", "standard-view-render"],
    })
    packaged["previewUrl"] = (packaged.get("viewImageUrls") or {}).get("front", "")
    if packaged.get("viewStatus") and packaged.get("viewStatus") != "ready":
        packaged["status"] = "needs_review"
        packaged["warnings"] = [packaged.get("viewError", "模型截图未通过检查")]
    (directory / "manifest.json").write_text(json.dumps(packaged, ensure_ascii=False), encoding="utf-8")
    return 200, "application/json; charset=utf-8", json.dumps(packaged, ensure_ascii=False).encode("utf-8")


def pack_glb(json_data, binary):
    json_bytes = json.dumps(json_data, separators=(",", ":"), ensure_ascii=False).encode("utf-8")
    json_bytes += b" " * ((4 - len(json_bytes) % 4) % 4)
    binary += b"\x00" * ((4 - len(binary) % 4) % 4)
    total_length = 12 + 8 + len(json_bytes) + 8 + len(binary)
    return b"glTF" + struct.pack("<II", 2, total_length) + struct.pack("<II", len(json_bytes), 0x4E4F534A) + json_bytes + struct.pack("<II", len(binary), 0x004E4942) + binary


def generate_demo_glb(body, content_type):
    fields, image = parse_multipart(body, content_type)
    source = Image.open(io.BytesIO(image["content"])).convert("RGBA")
    source.thumbnail((256, 256))
    texture_buffer = io.BytesIO()
    source.save(texture_buffer, format="PNG", optimize=True)
    texture = texture_buffer.getvalue()

    size = 32
    height_image = ImageOps.grayscale(source).resize((size, size), Image.Resampling.BILINEAR)
    pixels = list(height_image.getdata())
    positions = []
    uvs = []
    for z in range(size):
        for x in range(size):
            shade = pixels[z * size + x]
            height = (255 - shade) / 255 * 0.55
            positions.extend(((x / (size - 1) - 0.5) * 2.8, height, (z / (size - 1) - 0.5) * 2.8))
            uvs.extend((x / (size - 1), 1 - z / (size - 1)))
    indices = []
    for z in range(size - 1):
        for x in range(size - 1):
            a = z * size + x; b = a + 1; c = a + size; d = c + 1
            indices.extend((a, c, b, b, c, d))
    position_bytes = struct.pack(f"<{len(positions)}f", *positions)
    uv_bytes = struct.pack(f"<{len(uvs)}f", *uvs)
    index_bytes = struct.pack(f"<{len(indices)}I", *indices)
    binary = position_bytes + uv_bytes + index_bytes
    image_offset = len(binary)
    binary += texture
    min_y = min(positions[1::3]); max_y = max(positions[1::3])
    json_data = {
        "asset": {"version": "2.0", "generator": "Teacher Studio local demo"},
        "scene": 0, "scenes": [{"nodes": [0]}], "nodes": [{"mesh": 0, "name": "Local image relief"}],
        "meshes": [{"name": "Image relief", "primitives": [{"attributes": {"POSITION": 0, "TEXCOORD_0": 1}, "indices": 2, "material": 0}]}],
        "buffers": [{"byteLength": len(binary)}],
        "bufferViews": [
            {"buffer": 0, "byteOffset": 0, "byteLength": len(position_bytes), "target": 34962},
            {"buffer": 0, "byteOffset": len(position_bytes), "byteLength": len(uv_bytes), "target": 34962},
            {"buffer": 0, "byteOffset": len(position_bytes) + len(uv_bytes), "byteLength": len(index_bytes), "target": 34963},
            {"buffer": 0, "byteOffset": image_offset, "byteLength": len(texture)},
        ],
        "accessors": [
            {"bufferView": 0, "componentType": 5126, "count": size * size, "type": "VEC3", "min": [-1.4, min_y, -1.4], "max": [1.4, max_y, 1.4]},
            {"bufferView": 1, "componentType": 5126, "count": size * size, "type": "VEC2"},
            {"bufferView": 2, "componentType": 5125, "count": len(indices), "type": "SCALAR", "min": [0], "max": [size * size - 1]},
        ],
        "materials": [{"name": "Uploaded image", "pbrMetallicRoughness": {"baseColorTexture": {"index": 0}, "metallicFactor": 0, "roughnessFactor": 1}, "doubleSided": True}],
        "textures": [{"sampler": 0, "source": 0}], "samplers": [{"magFilter": 9729, "minFilter": 9987, "wrapS": 10497, "wrapT": 10497}],
        "images": [{"bufferView": 3, "mimeType": "image/png"}],
    }
    return pack_glb(json_data, binary)


def clean_lines(values, fallback):
    if isinstance(values, list):
        result = [str(v).strip() for v in values if str(v).strip()]
    else:
        result = [line.strip(" -•") for line in str(values or "").splitlines() if line.strip()]
    return result or [fallback]


def parse_material(body, content_type):
    message = BytesParser(policy=default).parsebytes(
        f"Content-Type: {content_type}\r\nMIME-Version: 1.0\r\n\r\n".encode() + body
    )
    if not message.is_multipart():
        raise ValueError("请上传教材文件")
    part = next((item for item in message.iter_parts() if item.get_param("name", header="content-disposition") == "material"), None)
    if part is None or not part.get_filename():
        raise ValueError("未找到教材文件")
    raw = part.get_payload(decode=True) or b""
    if not raw or len(raw) > 15 * 1024 * 1024:
        raise ValueError("文件不能为空且不能超过 15 MB")
    suffix = Path(part.get_filename()).suffix.lower()
    if suffix in (".txt", ".md"):
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = raw.decode("gb18030")
    elif suffix == ".docx":
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            xml = ET.fromstring(archive.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        content = "\n".join("".join(t.text or "" for t in p.findall(".//w:t", ns)) for p in xml.findall(".//w:p", ns))
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError("服务器尚未安装 pypdf，暂不能解析 PDF") from exc
        reader = PdfReader(io.BytesIO(raw))
        if len(reader.pages) > 100:
            raise ValueError("PDF 最多支持 100 页")
        content = "\n".join(page.extract_text() or "" for page in reader.pages)
    else:
        raise ValueError("仅支持 TXT、Markdown、DOCX 和文本型 PDF")
    content = content.strip()
    if not content:
        raise ValueError("未提取到文字；扫描版 PDF 请先进行 OCR")
    return content[:60000]


def visual_profile(request):
    benchmark = teaching_benchmarks.benchmark_profile(request)
    if request.get("stage") == "本科":
        return {"stage": "本科", "subject": "生物医学", "density": "balanced", "theme": "医学高对比", "visualLanguage": "器官标注、结构功能关系、机制箭头、病例证据链", "assetQueries": [], "benchmark": benchmark}
    subject = str(request.get("subject") or "").strip()
    stage = str(request.get("stage") or request.get("grade") or "").strip()
    if stage.startswith("小学") or "一" in stage or "二" in stage or "三" in stage or "四" in stage or "五" in stage or "六" in stage:
        density = "light"
    elif stage.startswith("高中") or "高" in stage:
        density = "dense"
    else:
        density = "balanced"
    rules = {
        "数学": {"visualLanguage": "坐标系、函数图像、几何关系、公式推导", "assetQueries": ["coordinate plane diagram", "geometry classroom illustration"]},
        "物理": {"visualLanguage": "受力图、电路图、实验步骤、变量关系", "assetQueries": ["physics force diagram", "school physics experiment"]},
        "化学": {"visualLanguage": "实验装置、反应过程、粒子示意、条件与结论", "assetQueries": ["chemistry laboratory equipment", "molecule structure diagram"]},
        "生物": {"visualLanguage": "结构标注、过程循环、对比图、实验观察", "assetQueries": ["biology cell diagram", "biology classroom microscope"]},
    }
    selected = rules.get(subject, {"visualLanguage": "问题链、流程图、对比卡片、课堂任务区域", "assetQueries": [f"{subject} classroom", f"{subject} learning diagram"]})
    theme = request.get("pptTheme") or "自动推荐"
    if theme == "自动推荐":
        theme = "考试复习" if "复习" in str(request.get("lessonType") or "") else "课堂故事化"
    return {"stage": stage, "subject": subject, "density": density, "theme": theme, **selected, "benchmark": benchmark}


def build_lesson(request):
    if request.get("stage") == "本科":
        return lesson_design.generate_draft(request, request.get("designMode") or "quick", request.get("modelAsset"))["lessonPlan"]
    subject = request.get("subject") or "未设置学科"
    grade = request.get("grade") or request.get("stage") or "未设置年级"
    topic = request.get("topic") or "待确定课题"
    duration = max(20, min(120, int(request.get("duration") or 40)))
    objectives = clean_lines(request.get("objectives"), "理解本节课的核心概念并完成迁移应用")
    key_points = clean_lines(request.get("keyPoints"), "建立概念、方法与实际问题之间的联系")
    difficulties = clean_lines(request.get("difficulties"), "把抽象知识转化为可操作的解题或探究步骤")
    methods = clean_lines(request.get("methods"), "问题驱动、合作探究、即时评价")
    situation = request.get("learningSituation") or "学生已有相关基础，需要通过真实问题建立新旧知识连接。"
    resources = request.get("resources") or "教材、课件、板书、课堂练习"
    assessment = request.get("assessment") or "观察参与度、口头表达、课堂练习和出口条"
    homework = request.get("homework") or "完成基础练习，并用一句话写下本节课最重要的结论。"
    lesson_type = request.get("lessonType") or "新授课"

    question = request.get("coreQuestion") or f"如何理解并运用{topic}？"
    prior = request.get("priorKnowledge") or "相关已有知识"
    errors = clean_lines(request.get("misconceptions"), "辨析常见误解")
    exam = clean_lines(request.get("examPoints"), "结合课堂目标设计检测题")
    material = (request.get("materialText") or "").strip()
    source = re.sub(r"\s+", " ", material[:400]).strip()
    rhythm = request.get("teachingRhythm") or "互动型"
    length = request.get("pptLength") or "标准"
    use_3d = request.get("use3d") != "不使用"
    times = [0.12, 0.22, 0.25, 0.27]
    allocations = [max(2, round(duration * part)) for part in times]
    allocations.append(duration - sum(allocations))
    if allocations[-1] < 2:
        allocations[3] -= 2 - allocations[-1]
        allocations[-1] = 2

    phases = [
        {"time": allocations[0], "name": "问题导入", "teacher": f"呈现核心问题：{question}；唤醒前置知识：{prior}。", "student": "提出初步判断，并说明已有依据。", "check": "记录至少一种初始想法"},
        {"time": allocations[1], "name": "观察与建构", "teacher": f"围绕{key_points[0]}组织观察、比较和归纳；引导学生回到资料核对。", "student": "比较实例，尝试用自己的语言概括规律。", "check": f"用一个例子解释：{key_points[0]}"},
        {"time": allocations[2], "name": "推理与示范", "teacher": f"示范从条件到结论的路径，重点突破：{difficulties[0]}。", "student": "独立完成关键一步，再与同伴互相解释。", "check": "能指出推理依据与适用条件"},
        {"time": allocations[3], "name": "辨析与迁移", "teacher": f"针对误区“{errors[0]}”组织反例讨论，再用题型“{exam[0]}”检查迁移。", "student": "先独立判断再交流，修正错误理由。", "check": "能说出错误在哪里及如何修正"},
        {"time": allocations[4], "name": "回扣与评价", "teacher": f"回到“{question}”，以出口条收集学习证据。", "student": "回答核心问题并写下仍未解决的疑问。", "check": assessment},
    ]
    profile = visual_profile(request)
    benchmark = profile["benchmark"]
    experience = benchmark["experience"]
    layout_by_kind = {
        "封面": "hero-stack", "问题": "question-split", "探索": "three-column", "教材": "source-quote",
        "概念": "layer-map", "方法": "step-flow", "辨析": "compare-duo", "练习": "task-timer",
        "揭示": "answer-rail", "考点": "exam-map", "观察": "observation-stage", "互动": experience["mode"], "总结": "loop-summary", "作业": "take-home",
    }
    proof_by_kind = {
        "封面": "visual-anchor", "问题": "question-prompt", "探索": "comparison-cards", "教材": "source-excerpt",
        "概念": "hierarchy-diagram", "方法": "step-sequence", "辨析": "claim-comparison", "练习": "answer-area",
        "揭示": "evidence-rail", "考点": "topic-map", "观察": "observation-task", "互动": "student-action", "总结": "question-loop", "作业": "task-list",
    }
    def page(kind, title, items, purpose, visual, interaction="", asset_query="", layout_override=""):
        return {
            "type": kind, "title": title, "items": items, "purpose": purpose, "visual": visual,
            "interaction": interaction, "layout": layout_override or layout_by_kind.get(kind, "content-list"),
            "proof": proof_by_kind.get(kind, "content-list"),
            "assetQuery": asset_query or (profile["assetQueries"][0] if kind in ("探索", "观察") else ""),
            "visualProfile": profile["visualLanguage"],
        }

    slides = [page("封面", topic, [f"{grade} · {subject} · {lesson_type}"], question, "核心问题作为视觉主角"),
              page("问题", "从一个问题开始", [question, f"你已经知道：{prior}"], "提出驱动全课的疑问", "大号问题与观察区域", "先独立猜想，再交流理由"),
              page("探索", "观察与发现", [key_points[0], f"尝试用已有知识解释：{prior}"], "从现象走向规律", "对比区与连接线", "说出相同与不同"),
              page("互动", experience["title"], experience["actions"], "让学生通过可见动作产生学习证据", "控制区、主画布与实时反馈", experience["feedback"], layout_override=experience["mode"]),
              page("概念", "建立核心认识", key_points, "明确概念与适用范围", "层级关系图"),
              page("方法", "怎样一步步解决", [difficulties[0], *objectives[:2]], "示范解决路径", "三步推理路线", "先补出缺失的一步"),
              page("辨析", "这个判断对吗？", [errors[0], "指出错误依据，并给出修正说法"], "暴露并修正常见误区", "双栏观点对照", "先投票后解释")]
    if source and request.get("materialPolicy") != "仅作为生成参考":
        slides.insert(3, page("教材", "回到教材", [source[:180]], "核对教材原文与课堂发现", "引文摘录"))
    if length == "详细":
        slides.append(page("练习", "先试一道基础题", [exam[0], "写出判断依据与适用条件"], "独立应用", "题干与作答区域", "先作答再交流"))
    slides.append(page("检测", "把方法用起来", exam[:2], "分层练习与迁移", "题干与作答区域", "独立完成后互评", layout_override="evidence-list"))
    if request.get("answerReveal") != "不包含":
        slides.append(page("揭示", "回看解题依据", [f"依据：{key_points[0]}", f"检查：{difficulties[0]}"], "练习后揭示思路", "逐项核对", "对照自己的过程修订"))
    if request.get("highlightExam") != "不突出" and exam[0] != "结合课堂目标设计检测题":
        slides.append(page("考点", "教师提供的考点", exam[:3], "聚焦教师填写的典型题型", "重点标记"))
    if use_3d and request.get("subject") in ("数学", "物理", "化学", "生物", "地理", "科学", "美术"):
        slides.append(page("观察", "换个视角观察", [f"观察任务：{question}", "准备模型后，比较不同角度下哪些特征不变"], "预留 3D 观察任务；模型需单独生成", "模型展示占位", "旋转模型并记录发现"))
    if rhythm == "练习型":
        slides.append(page("检测", "再试一个变化", [exam[-1], "解释与上一题的异同"], "加强迁移检测", "对比练习", "独立说明理由"))
    slides.append(page("总结", "回到最初的问题", [question, objectives[0], assessment], "形成性评价", "问题回环", "出口条回答"))
    if length != "精简":
        slides.append(page("作业", "带着问题继续", [homework], "巩固与延伸", "任务清单"))
    return {
        "id": uuid.uuid4().hex,
        "createdAt": now_iso(),
        "request": request,
        "title": topic,
        "overview": f"{grade}{subject}{lesson_type}，{duration} 分钟。核心问题：{question}。" + (f" 依据上传资料 {request.get('materialName') or '教师粘贴内容'}。" if source else " 依据教师填写内容。"),
        "objectives": objectives,
        "situation": situation,
        "keyPoints": key_points,
        "difficulties": difficulties,
        "methods": methods,
        "resources": resources,
        "assessment": assessment,
        "homework": homework,
        "phases": phases,
        "blackboard": [topic, key_points[0], "条件 → 方法 → 结论"],
        "slides": slides,
        "visualProfile": profile,
        "generationStandard": benchmark,
    }


def color(value):
    value = value.replace("#", "")
    return RGBColor.from_string(value)


def add_text(slide, value, x, y, w, h, size=18, fill="FFFFFF", bold=False, align=PP_ALIGN.LEFT):
    return courseware._add_text(slide, value, x, y, w, h, size, fill.replace('#', ''), bold, align)


def add_shape(slide, kind, x, y, w, h, fill, line_color=None, transparency=0):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = color(fill); s.fill.transparency = transparency
    if line_color:
        s.line.color.rgb = color(line_color); s.line.width = Pt(1)
    else:
        s.line.fill.background()
    return s


def expanded_slides(plan):
    if plan.get("courseProfile", {}).get("stage") == "本科":
        return [{**slide, "sourceIndex": i} for i, slide in enumerate(plan.get("slides", []))]
    result = []
    for source_index, item in enumerate(plan.get("slides", []) or []):
        values = [str(v).strip() for v in item.get("items", []) if str(v).strip()]
        chunks = [values[i:i + 3] for i in range(0, len(values), 3)] or [[]]
        for chunk in chunks:
            result.append({**item, "items": chunk, "sourceIndex": source_index})
    return result or [{"type": "封面", "title": plan.get("title", "课堂设计"), "items": [], "sourceIndex": 0}]


def quality_check(plan):
    if plan.get("courseProfile", {}).get("stage") == "本科":
        slide_quality = biomed_slides.quality(plan)
        design_quality = lesson_design.validate_lesson(plan, require_approval=True)
        blocking = [*design_quality["blocking"], *slide_quality["blocking"]]
        warnings = [*design_quality["warnings"], *slide_quality["warnings"]]
        return {**slide_quality, "blocking": blocking, "warnings": warnings, "score": max(0, 100 - len(blocking) * 25 - len(warnings) * 5), "designQuality": design_quality}
    slides = expanded_slides(plan)
    warnings = []
    blocking = []
    layouts = [slide.get("layout") or slide.get("type", "content-list") for slide in slides]
    if not slides:
        blocking.append({"code": "empty-deck", "message": "课件没有可生成的页面"})
    for index, slide in enumerate(slides, 1):
        title = str(slide.get("title") or "").strip()
        items = [str(value).strip() for value in slide.get("items", []) if str(value).strip()]
        if not title:
            blocking.append({"code": "empty-title", "slide": index, "message": "页面缺少标题"})
        if len(title) > 48:
            warnings.append({"code": "long-title", "slide": index, "message": "标题较长，可能需要人工压缩"})
        if sum(len(item) for item in items) > 420:
            warnings.append({"code": "dense-text", "slide": index, "message": "页面文字较多，建议拆页或删减"})
        if not slide.get("proof"):
            warnings.append({"code": "missing-proof", "slide": index, "message": "页面缺少证明对象元数据"})
    for index in range(1, len(layouts)):
        if layouts[index] == layouts[index - 1]:
            warnings.append({"code": "repeated-layout", "slide": index + 1, "message": "连续页面使用了相同版式"})
    types = {slide.get("type") for slide in slides}
    for required in ("问题", "总结"):
        if required not in types:
            warnings.append({"code": f"missing-{required}", "message": f"课件缺少{required}页"})
    if not ({"练习", "检测"} & types):
        warnings.append({"code": "missing-practice", "message": "课件缺少练习或检测页"})
    if not any(slide.get("type") in ("互动", "观察") for slide in slides):
        warnings.append({"code": "missing-interaction", "message": "课件缺少学生可操作的实验、角色或挑战任务"})
    if not any(slide.get("type") in ("检测", "揭示") for slide in slides):
        warnings.append({"code": "missing-feedback", "message": "课件缺少即时检测或答案依据反馈"})
    if not any(slide.get("type") in ("检测", "作业") for slide in slides):
        warnings.append({"code": "missing-transfer", "message": "课件缺少迁移任务"})
    score = max(0, min(100, 100 - len(blocking) * 35 - len(warnings) * 5))
    return {"score": score, "blocking": blocking, "warnings": warnings, "slideCount": len(slides), "layouts": sorted(set(layouts))}


def preview_theme(plan):
    request = plan.get("request") or {}
    theme = request.get("pptTheme", "自动推荐")
    if theme == "自动推荐":
        theme = "考试复习" if "复习" in request.get("lessonType", "") else "课堂故事化"
    return {
        "课堂故事化": ("080B0A", "F4F7F2", "A4FF02", "1A1D1A", "A8B0AA"),
        "考试复习": ("111315", "FFF4EA", "FF7B55", "25282A", "B9AAA1"),
        "科技展示": ("0B0D10", "FFF5EA", "FFA649", "20252B", "BBB5AE"),
    }.get(theme, ("080B0A", "F4F7F2", "A4FF02", "1A1D1A", "A8B0AA"))


def svg_text(value, x, y, size, fill, weight="400", max_chars=34):
    value = html.escape(re.sub(r"\s+", " ", str(value or "")))
    lines = [value[i:i + max_chars] for i in range(0, len(value), max_chars)] or [""]
    return "".join(f'<text x="{x}" y="{y + index * (size + 7)}" fill="#{fill}" font-size="{size}" font-family="Microsoft YaHei, sans-serif" font-weight="{weight}">{line}</text>' for index, line in enumerate(lines[:4]))


def render_preview_svg(slide, index, plan):
    if plan.get("courseProfile", {}).get("stage") == "本科":
        return biomed_slides.render_svg(slide, index, plan)
    bg, ink, accent, surface, muted = preview_theme(plan)
    title = slide.get("title", "课堂设计")
    items = slide.get("items", []) or []
    kind = slide.get("type", "教学页")
    layout = slide.get("layout", "content-list")
    parts = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720" role="img" aria-label="{html.escape(title)}">', f'<rect width="1280" height="720" fill="#{bg}"/>']
    for x in range(54, 1280, 102):
        parts.append(f'<path d="M{x} 0V720" stroke="#{surface}" stroke-width="1" opacity=".5"/>')
    parts.append(f'<rect width="10" height="720" fill="#{accent}"/>')
    parts.append(svg_text(f"课研台 / {kind}", 68, 38, 13, accent, "700", 32))
    parts.append(svg_text(f"{index:02d}", 1180, 38, 17, accent, "700", 4))
    parts.append(svg_text(title, 72, 112, 34, ink, "700", 30))
    parts.append(f'<rect x="72" y="188" width="64" height="5" fill="#{accent}"/>')
    if layout in ("hero-stack", "question-split"):
        parts.append(svg_text(items[0] if items else slide.get("purpose", ""), 76, 270, 29, ink, "700", 31))
        parts.append(f'<rect x="820" y="220" width="360" height="250" rx="22" fill="#{surface}" stroke="#{accent}" stroke-width="2"/>')
        parts.append(svg_text("课堂任务", 860, 275, 18, accent, "700", 12))
        parts.append(svg_text(items[1] if len(items) > 1 else slide.get("interaction", ""), 860, 335, 20, muted, "400", 18))
    elif layout in ("three-column", "step-flow"):
        for n, value in enumerate(items[:3]):
            x = 72 + n * 390
            parts.append(f'<rect x="{x}" y="250" width="320" height="250" rx="18" fill="#{surface}" stroke="#{accent}" stroke-width="2"/>')
            parts.append(svg_text(f"0{n + 1}", x + 24, 290, 27, accent, "700", 5))
            parts.append(svg_text(value, x + 24, 355, 21, ink, "700", 19))
    elif layout in ("control-lab", "roleplay-stage", "game-challenge"):
        label = {"control-lab": "CONTROL LAB", "roleplay-stage": "ROLE PLAY", "game-challenge": "CHALLENGE"}[layout]
        parts.append(f'<rect x="72" y="230" width="300" height="330" rx="14" fill="#{surface}" stroke="#{accent}" stroke-width="2"/>')
        parts.append(svg_text(label, 104, 278, 18, accent, "700", 18))
        for n, value in enumerate(items[:3]):
            parts.append(f'<rect x="104" y="{318 + n * 70}" width="230" height="8" rx="4" fill="#{accent}" opacity="{1 - n * .22}"/>')
            parts.append(svg_text(f"0{n + 1}", 104, 354 + n * 70, 14, muted, "700", 5))
        parts.append(f'<rect x="420" y="230" width="780" height="330" rx="14" fill="#{surface}" stroke="#{accent}" stroke-width="2"/>')
        parts.append(f'<circle cx="810" cy="365" r="92" fill="none" stroke="#{accent}" stroke-width="5"/>')
        parts.append(f'<path d="M710 430L770 380L824 405L914 290" fill="none" stroke="#{accent}" stroke-width="7"/>')
        parts.append(svg_text(slide.get("interaction", "操作后立即核对证据"), 500, 510, 18, ink, "700", 32))
    elif layout in ("compare-duo", "source-quote"):
        for n in range(2):
            x = 72 + n * 600
            parts.append(f'<rect x="{x}" y="235" width="550" height="285" rx="18" fill="#{accent if n == 1 and layout == "source-quote" else surface}" stroke="#{accent}" stroke-width="2"/>')
            parts.append(svg_text("资料摘录" if layout == "source-quote" and n == 1 else ("判断" if n == 0 else "修正"), x + 28, 280, 17, bg if n == 1 and layout == "source-quote" else accent, "700", 12))
            parts.append(svg_text(items[n] if n < len(items) else slide.get("purpose", ""), x + 28, 350, 22, bg if n == 1 and layout == "source-quote" else ink, "700", 23))
    elif layout in ("layer-map", "exam-map"):
        for n, value in enumerate(items[:3]):
            x = 84 + n * 210
            y = 260 + n * 72
            parts.append(f'<rect x="{x}" y="{y}" width="{560 - n * 80}" height="54" rx="9" fill="#{accent if n == 0 else surface}"/>')
            parts.append(svg_text(value, x + 20, y + 34, 16, bg if n == 0 else ink, "700", 42))
        parts.append(svg_text("关系 / 结构 / 证据", 760, 330, 24, accent, "700", 18))
    elif layout == "evidence-list":
        parts.append(f'<rect x="72" y="235" width="555" height="285" rx="18" fill="#{surface}" stroke="#{accent}" stroke-width="2"/>')
        parts.append(f'<rect x="680" y="235" width="520" height="285" rx="18" fill="#{surface}" stroke="#{accent}" stroke-width="2"/>')
        parts.append(svg_text(items[0] if items else slide.get("purpose", ""), 104, 315, 22, ink, "700", 24))
        parts.append(svg_text("检查清单", 720, 300, 18, accent, "700", 14))
        for n, value in enumerate(["条件是否对应？", "方法是否适用？", "结论能否解释？"]):
            parts.append(svg_text(f"0{n + 1}", 720, 365 + n * 48, 15, accent, "700", 5))
            parts.append(svg_text(value, 770, 365 + n * 48, 16, ink, "700", 20))
    else:
        for n, value in enumerate(items[:3]):
            y = 265 + n * 92
            parts.append(svg_text(f"0{n + 1}", 76, y, 22, accent, "700", 5))
            parts.append(svg_text(value, 150, y, 20, ink, "400", 52))
            parts.append(f'<path d="M76 {y + 35}H1200" stroke="#{surface}" stroke-width="2"/>')
    parts.append(svg_text("课研台 · 可编辑课件", 72, 695, 11, muted, "400", 28))
    parts.append("</svg>")
    return "".join(parts)


def search_assets(query, limit=6):
    query = str(query or "").strip()
    if not query:
        return []
    url = "https://api.openverse.org/v1/images/?q=" + quote(query) + f"&page_size={max(1, min(int(limit), 12))}"
    try:
        request = Request(url, headers={"Accept": "application/json", "User-Agent": "TeacherStudio/1.0"})
        with urlopen(request, timeout=8) as response:
            payload = json.loads(response.read().decode("utf-8"))
        return [{"sourceType": "openverse", "localPath": "", "sourceUrl": item.get("url", ""), "license": item.get("license", "unknown"), "altText": item.get("title") or query, "cropMode": "cover", "subjectTags": [query], "thumbnail": item.get("thumbnail", "")} for item in payload.get("results", []) if item.get("url")][:12]
    except (OSError, ValueError, HTTPError, URLError):
        return []


def cache_asset(url, metadata=None):
    if not url or urlparse(url).scheme not in ("http", "https"):
        raise ValueError("素材地址必须是 http 或 https")
    request = Request(url, headers={"User-Agent": "TeacherStudio/1.0"})
    with urlopen(request, timeout=12) as response:
        content_type = response.headers.get_content_type()
        if content_type not in ("image/png", "image/jpeg", "image/webp", "image/gif"):
            raise ValueError("只允许缓存图片素材")
        data = response.read(8 * 1024 * 1024 + 1)
    if len(data) > 8 * 1024 * 1024:
        raise ValueError("图片不能超过 8 MB")
    suffix = {"image/png": ".png", "image/jpeg": ".jpg", "image/webp": ".webp", "image/gif": ".gif"}[content_type]
    name = hashlib.sha256(url.encode("utf-8")).hexdigest()[:20] + suffix
    path = ASSET_DIR / name
    path.write_bytes(data)
    record = {"sourceType": "cached", "localPath": str(path.relative_to(ROOT)), "sourceUrl": url, **(metadata or {})}
    (path.with_suffix(path.suffix + ".json")).write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
    return record


def add_footer(slide, index):
    add_text(slide, "课研台 · 教师备案工作台", 0.55, 7.03, 3.0, 0.18, 8, "7090B3")
    add_text(slide, f"{index:02d}", 12.20, 7.00, 0.55, 0.20, 9, "7AD8FF", True, PP_ALIGN.RIGHT)


def create_ppt(plan):
    """Build a varied, dark editorial classroom deck inspired by user references."""
    if plan.get("courseProfile", {}).get("stage") == "本科":
        return biomed_slides.create_ppt(plan, OUTPUT_DIR)
    prs = Presentation(); prs.slide_width = Inches(13.333333); prs.slide_height = Inches(7.5)
    request = plan.get("request") or {}
    theme = request.get("pptTheme", "自动推荐")
    if theme == "自动推荐":
        theme = "考试复习" if "复习" in request.get("lessonType", "") else "课堂故事化"
    palettes = {
        "课堂故事化": ("080B0A", "F4F7F2", "A4FF02", "1A1D1A", "A8B0AA", "343B34"),
        "考试复习": ("111315", "FFF4EA", "FF7B55", "25282A", "B9AAA1", "484044"),
        "科技展示": ("0B0D10", "FFF5EA", "FFA649", "20252B", "BBB5AE", "45494D"),
    }
    bg, ink, accent, surface, muted, line = palettes.get(theme, palettes["课堂故事化"])
    slides = [(item, item.get("items", [])) for item in expanded_slides(plan)]

    def short(value, limit=120):
        value = re.sub(r"\s+", " ", str(value or ""))
        return value if len(value) <= limit else value[:limit - 1] + "…"

    def panel(slide, x, y, w, h, fill=surface, border=line, rounded=True):
        return add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h, fill, border)

    def chip(slide, value, x, y, w=1.05, fill=accent, text_fill=bg):
        panel(slide, x, y, w, 0.30, fill, fill)
        add_text(slide, value, x + 0.04, y + 0.045, w - 0.08, 0.18, 9, text_fill, True, PP_ALIGN.CENTER)

    def chrome(slide, kind, index):
        for i in range(13):
            add_shape(slide, MSO_SHAPE.RECTANGLE, 0.55 + i * 1.02, 0, 0.006, 7.5, line, transparency=66)
        for i in range(10):
            add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0.62 + i * 0.72, 13.333, 0.006, line, transparency=72)
        add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.10, 7.5, accent)
        add_text(slide, f"课研台  /  {kind}", 0.70, 0.27, 4.5, 0.22, 9, accent, True)
        add_text(slide, "TEACHER STUDIO", 10.30, 0.27, 1.8, 0.22, 8, muted, True, PP_ALIGN.RIGHT)
        add_text(slide, f"{index:02d}", 12.18, 0.20, 0.52, 0.30, 13, accent, True, PP_ALIGN.RIGHT)

    def title_block(slide, title, subtitle="", width=9.2):
        add_text(slide, short(title, 42), 0.78, 0.90, width, 0.60, 27, ink, True)
        if subtitle:
            add_text(slide, short(subtitle, 115), 0.78, 1.54, width, 0.38, 11, muted)
        add_shape(slide, MSO_SHAPE.RECTANGLE, 0.78, 2.02, 0.72, 0.05, accent)

    def footer(slide):
        add_text(slide, short(plan.get("title", "课堂设计"), 34), 0.78, 7.04, 6.8, 0.18, 8, muted)
        add_text(slide, "可编辑课件 · 课堂证据优先", 9.2, 7.04, 3.3, 0.18, 8, muted, False, PP_ALIGN.RIGHT)

    for index, (item, items) in enumerate(slides, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = color(bg)
        original_kind = item.get("type", "教学页"); title = item.get("title", "课堂设计")
        layout_kind = {"hero-stack": "封面", "question-split": "问题", "three-column": "探索", "source-quote": "教材", "layer-map": "概念", "step-flow": "方法", "compare-duo": "辨析", "task-timer": "练习", "evidence-list": "检测", "answer-rail": "揭示", "exam-map": "考点", "observation-stage": "观察", "control-lab": "互动实验", "roleplay-stage": "互动实验", "game-challenge": "互动实验", "loop-summary": "总结", "take-home": "作业"}
        kind = layout_kind.get(item.get("layout"), original_kind)
        purpose = item.get("purpose", ""); lead = items[0] if items else purpose
        chrome(slide, original_kind, index)

        if kind == "封面":
            add_text(slide, short(title, 30), 0.78, 1.50, 7.0, 1.12, 36, ink, True)
            add_text(slide, short(purpose or plan.get("overview", ""), 105), 0.82, 2.95, 6.6, 1.05, 16, muted)
            chip(slide, "LESSON / STORY", 0.82, 4.48, 1.48)
            add_text(slide, short(" · ".join(items), 110), 0.82, 5.10, 6.8, 0.44, 14, ink)
            for x, y, w, h, fill in [(8.10, 1.20, 3.75, 3.3, surface), (8.95, 1.95, 3.45, 3.3, line), (8.55, 3.00, 3.05, 2.72, accent)]:
                shape = panel(slide, x, y, w, h, fill, fill, False); shape.rotation = -8
            add_text(slide, "01", 9.10, 3.38, 0.8, 0.42, 24, bg, True)
            add_text(slide, "THINK\nEXPLORE\nEXPLAIN", 9.10, 4.05, 2.15, 1.05, 15, bg, True)
        elif kind == "问题":
            title_block(slide, title, "把第一反应变成可验证的学习任务。")
            add_text(slide, short(lead, 92), 0.82, 2.50, 7.05, 1.30, 28, ink, True)
            panel(slide, 8.45, 2.20, 3.95, 3.28)
            chip(slide, "先猜想", 8.82, 2.58, 1.05)
            add_text(slide, "你已经知道什么？", 8.82, 3.20, 3.1, 0.34, 17, ink, True)
            add_text(slide, short(items[1] if len(items) > 1 else "写下一个依据，再和同伴比较。", 90), 8.82, 3.82, 3.05, 0.95, 14, muted)
            add_shape(slide, MSO_SHAPE.RECTANGLE, 8.82, 5.02, 2.55, 0.04, accent)
        elif kind == "探索":
            title_block(slide, title, "让发现先发生，再补上准确术语。")
            for n, content in enumerate((items or [lead])[:3]):
                x = 0.82 + n * 4.05
                panel(slide, x, 2.40, 3.45, 3.00)
                add_text(slide, f"0{n + 1}", x + 0.24, 2.68, 0.65, 0.42, 23, accent, True)
                add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.24, 3.35, 2.8, 0.04, accent)
                add_text(slide, short(content, 70), x + 0.24, 3.75, 2.85, 1.14, 16, ink, True)
                add_text(slide, ["观察", "比较", "表达"][n], x + 0.24, 5.03, 1.4, 0.22, 10, muted)
        elif kind == "互动实验":
            mode = item.get("layout", "control-lab")
            title_block(slide, title, "先产生预测，再用操作结果修正解释。")
            panel(slide, 0.82, 2.30, 3.25, 3.35)
            chip(slide, {"control-lab": "CONTROL", "roleplay-stage": "ROLE", "game-challenge": "CHALLENGE"}.get(mode, "ACTION"), 1.14, 2.62, 1.25)
            for n, content in enumerate((items or [lead])[:3]):
                y = 3.20 + n * 0.70
                add_text(slide, f"0{n + 1}", 1.14, y, 0.42, 0.28, 13, accent, True)
                add_text(slide, short(content, 42), 1.72, y - 0.03, 1.95, 0.52, 11, ink, True)
            panel(slide, 4.48, 2.30, 7.65, 3.35)
            add_text(slide, "LIVE LEARNING CANVAS", 4.84, 2.62, 3.2, 0.24, 10, accent, True)
            for x, y, size in [(6.08, 3.28, 1.48), (7.64, 3.04, 1.05), (8.86, 3.55, 0.68)]:
                shape = add_shape(slide, MSO_SHAPE.OVAL, x, y, size, size, surface, accent); shape.line.width = Pt(2)
            add_shape(slide, MSO_SHAPE.CHEVRON, 7.18, 4.52, 1.18, 0.45, accent)
            add_text(slide, short(item.get("interaction") or "观察变化，说明证据。", 88), 5.02, 5.02, 6.45, 0.36, 12, ink, True)
        elif kind == "教材":
            title_block(slide, title, "把资料原话变成课堂证据。")
            panel(slide, 0.82, 2.25, 7.35, 3.40)
            add_text(slide, "SOURCE EXCERPT", 1.18, 2.62, 2.3, 0.22, 9, accent, True)
            add_text(slide, f'“{short(lead, 250)}”', 1.18, 3.16, 6.45, 1.62, 20, ink, True)
            panel(slide, 8.62, 2.25, 3.5, 3.40, accent, accent)
            add_text(slide, "课堂转化", 8.96, 2.62, 2.5, 0.28, 12, bg, True)
            add_text(slide, short(purpose, 82), 8.96, 3.22, 2.72, 1.16, 16, bg, True)
            add_text(slide, "来源：教师上传资料", 8.96, 5.08, 2.5, 0.22, 9, bg)
        elif kind == "概念":
            title_block(slide, title, "用层级和关系建立可复述的结构模型。")
            subject = str(request.get("subject") or "")
            labels = {
                "数学": ("定义域", "变化趋势", "单调区间"),
                "物理": ("条件", "运动过程", "物理结论"),
                "化学": ("反应条件", "微观过程", "宏观现象"),
                "生物": ("结构", "功能过程", "实验结论"),
            }.get(subject, ("核心概念", "关键关系", "可验证结论"))
            for label, x, y, w, fill in [(labels[0], 0.88, 2.35, 4.9, accent), (labels[1], 1.48, 3.30, 4.3, surface), (labels[2], 2.08, 4.25, 3.7, line)]:
                panel(slide, x, y, w, 0.65, fill, fill)
                add_text(slide, label, x + 0.20, y + 0.18, w - 0.4, 0.24, 12, bg if fill == accent else ink, True)
            if subject == "数学":
                panel(slide, 6.65, 2.35, 5.25, 3.05, surface, line, False)
                add_shape(slide, MSO_SHAPE.RECTANGLE, 7.05, 4.82, 4.25, 0.025, muted)
                add_shape(slide, MSO_SHAPE.RECTANGLE, 7.35, 2.72, 0.025, 2.10, muted)
                for tick in range(1, 5):
                    add_shape(slide, MSO_SHAPE.RECTANGLE, 7.35 + tick * 0.76, 4.74, 0.018, 0.18, muted)
                for tick in range(1, 4):
                    add_shape(slide, MSO_SHAPE.RECTANGLE, 7.27, 4.82 - tick * 0.52, 0.18, 0.018, muted)
                for x, y, w, h, angle in [(7.55, 4.28, 1.02, 0.05, -24), (8.50, 3.86, 1.04, 0.05, -24), (9.48, 3.44, 1.04, 0.05, -24)]:
                    graph = add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, accent)
                    graph.rotation = angle
                add_text(slide, "y", 7.05, 2.48, 0.24, 0.22, 12, accent, True)
                add_text(slide, "x", 11.26, 4.88, 0.24, 0.22, 12, accent, True)
                add_text(slide, "图像上升  →  函数值增大", 7.55, 5.08, 3.95, 0.24, 11, ink, True)
            else:
                add_text(slide, "核心关系", 6.65, 2.45, 1.4, 0.22, 10, accent, True)
                add_shape(slide, MSO_SHAPE.RECTANGLE, 6.65, 2.84, 4.9, 0.05, accent)
                for n, content in enumerate(items[:3]):
                    add_text(slide, f"0{n + 1}", 6.65, 3.20 + n * 0.83, 0.45, 0.35, 15, accent, True)
                    add_text(slide, short(content, 74), 7.28, 3.16 + n * 0.83, 4.2, 0.52, 13, ink)
        elif kind == "方法":
            title_block(slide, title, "一条可检查、可复述的推理路线。")
            for n, content in enumerate((items or [lead])[:3]):
                x = 0.82 + n * 4.05
                panel(slide, x, 2.38, 3.35, 2.72)
                chip(slide, f"STEP 0{n + 1}", x + 0.24, 2.70, 1.05)
                add_text(slide, short(content, 64), x + 0.24, 3.52, 2.78, 1.02, 18, ink, True)
                if n < 2: add_shape(slide, MSO_SHAPE.CHEVRON, x + 3.52, 3.42, 0.42, 0.42, accent)
            add_text(slide, "先独立完成关键一步，再交换依据。", 0.84, 5.52, 6.8, 0.28, 12, muted)
        elif kind == "辨析":
            title_block(slide, title, "把错误说清楚，比只给正确答案更重要。")
            for n, content in enumerate((items or [lead])[:2]):
                x = 0.82 + n * 6.05
                panel(slide, x, 2.34, 5.45, 3.10)
                chip(slide, "常见判断" if n == 0 else "修正说法", x + 0.28, 2.68, 1.15, accent if n == 0 else ink, bg if n == 0 else surface)
                add_text(slide, short(content, 105), x + 0.28, 3.42, 4.65, 1.18, 20, ink, True)
                add_text(slide, "请先投票，再说依据。", x + 0.28, 4.91, 3.6, 0.24, 10, muted)
        elif kind == "练习":
            title_block(slide, title, "先留下思考痕迹，再讨论答案。")
            panel(slide, 0.82, 2.32, 7.35, 3.38); chip(slide, "TASK", 1.16, 2.68, 0.75)
            add_text(slide, short(lead, 150), 1.16, 3.26, 6.45, 1.22, 22, ink, True)
            add_text(slide, "我的依据：____________________________", 1.16, 5.08, 5.9, 0.25, 11, muted)
            panel(slide, 8.62, 2.32, 3.5, 3.38, accent, accent)
            add_text(slide, "3 MIN", 8.98, 2.78, 2.3, 0.58, 28, bg, True)
            add_text(slide, "独立完成\n再与同伴互评", 8.98, 3.75, 2.55, 0.75, 17, bg, True)
        elif kind == "检测":
            title_block(slide, title, "把方法迁移到新的题型或情境。")
            panel(slide, 0.82, 2.32, 5.55, 3.35)
            chip(slide, "TRANSFER", 1.16, 2.68, 1.15)
            add_text(slide, short(lead, 105), 1.16, 3.32, 4.75, 1.15, 20, ink, True)
            add_text(slide, "我采用的方法：________________", 1.16, 5.08, 4.5, 0.24, 11, muted)
            panel(slide, 6.82, 2.32, 5.3, 3.35)
            add_text(slide, "检查清单", 7.20, 2.70, 2.0, 0.26, 12, accent, True)
            for n, label in enumerate(["条件是否对应？", "方法是否适用？", "结论能否解释？"]):
                y = 3.30 + n * 0.68
                add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 7.20, y, 0.28, 0.28, accent)
                add_text(slide, label, 7.70, y + 0.02, 3.7, 0.24, 14, ink, True)
        elif kind == "揭示":
            title_block(slide, title, "答案不是终点，要回看它依赖了什么。")
            add_text(slide, "ANSWER", 0.84, 2.45, 2.0, 0.32, 12, accent, True)
            add_text(slide, short(lead, 88), 0.84, 3.00, 6.75, 1.22, 28, ink, True)
            for n, content in enumerate(items[1:3]):
                panel(slide, 8.25, 2.34 + n * 1.55, 4.0, 1.18)
                add_text(slide, f"0{n + 1}", 8.58, 2.68 + n * 1.55, 0.45, 0.25, 14, accent, True)
                add_text(slide, short(content, 54), 9.25, 2.58 + n * 1.55, 2.60, 0.48, 12, ink, True)
        elif kind == "考点":
            title_block(slide, title, "把考点转成一眼能扫到的复习地图。")
            for n, content in enumerate((items or [lead])[:3]):
                x = 0.82 + n * 2.05; fill = accent if n == 0 else surface
                panel(slide, x, 2.38, 1.72, 2.08, fill, accent if n == 0 else line)
                add_text(slide, f"0{n + 1}", x + 0.18, 2.64, 0.5, 0.28, 15, bg if n == 0 else accent, True)
                add_text(slide, short(content, 38), x + 0.18, 3.30, 1.34, 0.72, 13, bg if n == 0 else ink, True)
            panel(slide, 7.35, 2.38, 4.75, 2.08)
            add_text(slide, "高频动作", 7.72, 2.72, 2.0, 0.25, 11, accent, True)
            add_text(slide, "定位条件  →  选择方法  →  检查结论", 7.72, 3.40, 3.85, 0.34, 16, ink, True)
        elif kind == "观察":
            title_block(slide, title, "把抽象结构换成可旋转、可比较的观察任务。")
            panel(slide, 0.86, 2.30, 5.0, 3.32)
            for n, size in enumerate((2.35, 1.55, 0.82)):
                shape = panel(slide, 2.18 + (2.35-size)/2, 2.78 + (2.35-size)/2, size, size, accent if n == 2 else line, accent if n == 2 else line, False); shape.rotation = 12
            add_text(slide, "ROTATE / COMPARE", 1.30, 5.12, 3.8, 0.28, 11, accent, True, PP_ALIGN.CENTER)
            panel(slide, 6.45, 2.30, 5.65, 3.32)
            add_text(slide, short(lead, 84), 6.86, 2.82, 4.8, 0.92, 19, ink, True)
            add_text(slide, short(items[1] if len(items) > 1 else "记录观察到的不变量。", 90), 6.86, 4.18, 4.55, 0.70, 13, muted)
        elif kind == "总结":
            title_block(slide, title, "让学生带着最初的问题离开课堂。")
            for n, label in enumerate(["问题", "证据", "结论"]):
                x = 1.0 + n * 3.55; fill = accent if n == 2 else surface
                panel(slide, x, 2.60, 2.35, 1.28, fill, accent if n == 2 else line)
                add_text(slide, f"0{n + 1}", x + 0.18, 2.88, 0.42, 0.25, 14, bg if n == 2 else accent, True)
                add_text(slide, label, x + 0.70, 2.86, 1.25, 0.28, 17, bg if n == 2 else ink, True)
                if n < 2: add_shape(slide, MSO_SHAPE.CHEVRON, x + 2.62, 2.98, 0.48, 0.45, accent)
            add_text(slide, short(lead, 105), 1.0, 4.62, 10.9, 0.72, 22, ink, True, PP_ALIGN.CENTER)
        elif kind == "作业":
            title_block(slide, title, "把课堂结构带回一个真实文件。")
            panel(slide, 0.82, 2.35, 7.15, 3.28, accent, accent)
            add_text(slide, "TAKE HOME", 1.20, 2.75, 2.3, 0.24, 10, bg, True)
            add_text(slide, short(lead, 115), 1.20, 3.42, 6.05, 1.16, 21, bg, True)
            for n, label in enumerate(["复制副本", "解压观察", "记录关系"]):
                y = 2.38 + n * 0.96
                panel(slide, 8.55, y, 3.45, 0.72)
                add_text(slide, f"0{n + 1}", 8.77, y + 0.20, 0.35, 0.22, 13, accent, True)
                add_text(slide, label, 9.30, y + 0.18, 2.2, 0.25, 14, ink, True)
        else:
            title_block(slide, title, purpose)
            for n, content in enumerate(items[:3]):
                y = 2.38 + n * 1.10
                add_text(slide, f"0{n + 1}", 0.82, y, 0.75, 0.42, 18, accent, True)
                add_text(slide, short(content, 125), 1.65, y - 0.04, 10.7, 0.84, 18, ink)
                add_shape(slide, MSO_SHAPE.RECTANGLE, 0.82, y + 0.88, 11.55, 0.012, line)

        if request.get("speakerNotes") != "不包含":
            slide.notes_slide.notes_text_frame.text = f"教学意图：{purpose}\n视觉建议：{item.get('visual', '')}\n互动：{item.get('interaction', '')}"
        footer(slide)
    output = OUTPUT_DIR / f"{re.sub(r'[^a-zA-Z0-9_-]+', '_', plan.get('title', 'lesson'))}_{uuid.uuid4().hex[:8]}.pptx"
    prs.save(output)
    return output


class Handler(BaseHTTPRequestHandler):
    def _headers(self, content_type="application/json; charset=utf-8", length=None):
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET,POST,OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.send_header("Content-Type", content_type)
        if length is not None:
            self.send_header("Content-Length", str(length))

    def do_OPTIONS(self):
        self.send_response(204); self._headers(); self.end_headers()

    def send_json(self, payload, status=200):
        data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status); self._headers(length=len(data)); self.end_headers(); self.wfile.write(data)

    def send_bytes(self, data, content_type, status=200):
        self.send_response(status); self._headers(content_type, len(data)); self.end_headers(); self.wfile.write(data)

    def do_POST(self):
        try:
            request_path = urlparse(self.path).path
            length = int(self.headers.get("Content-Length", "0"))
            max_bytes = 75 * 1024 * 1024 if request_path.startswith("/api/3d/") else 20 * 1024 * 1024
            if length <= 0 or length > max_bytes:
                self.send_json({"error": f"请求不能为空且不能超过 {max_bytes // 1024 // 1024} MB"}, 413)
                return
            body = self.rfile.read(length)
            if request_path == "/api/material/parse":
                content = parse_material(body, self.headers.get("Content-Type", ""))
                self.send_json({"text": content, "characters": len(content)})
                return
            if request_path == "/api/3d/generate":
                status, content_type, result = generate_image_model(body, self.headers.get("Content-Type", ""), "single")
                self.send_bytes(result, content_type, status)
                return
            if request_path == "/api/3d/generate-image":
                status, content_type, result = generate_image_model(body, self.headers.get("Content-Type", ""), "single")
                self.send_bytes(result, content_type, status)
                return
            if request_path == "/api/3d/generate-multiview":
                status, content_type, result = generate_image_model(body, self.headers.get("Content-Type", ""), "multiview")
                self.send_bytes(result, content_type, status)
                return
            if request_path == "/api/3d/generate-multiview-model":
                status, content_type, result = generate_image_model(body, self.headers.get("Content-Type", ""), "multiview")
                self.send_bytes(result, content_type, status)
                return
            if request_path == "/api/3d/demo":
                raise RuntimeError("本地演示建模已停用，请使用 混元 3D 生成模型")
            if request_path == "/api/3d/multiview-demo":
                raise RuntimeError("本地演示建模已停用，请使用 混元 3D 生成模型")
            payload = json.loads(body or b"{}")
            if request_path == "/api/warmup/generate":
                items, generation, warnings = generate_warmup_with_ai(payload, payload.get("currentQuestions"))
                lesson = payload.get("lessonPlan") or {}
                model_url, _parts = _warmup_model_context(payload.get("model"))
                questions = _warmup_with_metadata(items, lesson, model_url, generation)
                self.send_json({"questions": questions, "count": 5, "generation": generation, "warnings": warnings})
            elif request_path == "/api/warmup/regenerate-question":
                index = max(0, min(4, int(payload.get("index", 0))))
                items, generation, warnings = generate_warmup_with_ai(payload, payload.get("currentQuestions"), index)
                lesson = payload.get("lessonPlan") or {}
                model_url, _parts = _warmup_model_context(payload.get("model"))
                questions = _warmup_with_metadata(items, lesson, model_url, generation)
                self.send_json({"question": questions[index], "index": index, "generation": generation, "warnings": warnings})
            elif request_path == "/api/lesson-design/generate-draft":
                self.send_json(lesson_design.generate_draft(payload.get("request") or payload, payload.get("mode") or "quick", payload.get("model")))
            elif request_path == "/api/lesson-design/generate-step":
                self.send_json(lesson_design.generate_step(payload))
            elif request_path == "/api/lesson-design/edit-step":
                self.send_json(lesson_design.edit_step(payload))
            elif request_path == "/api/lesson-design/sync-downstream":
                self.send_json(lesson_design.sync_downstream(payload))
            elif request_path == "/api/lesson-design/validate":
                self.send_json(lesson_design.validate_payload(payload))
            elif request_path == "/api/lesson-design/prepare-export":
                prepared = lesson_design.prepare_teacher_lesson(payload)
                status_code = 200 if prepared.get("draft") else 503
                self.send_json(prepared, status_code)
            elif request_path == "/api/courseware/generate":
                lesson_plan = lesson_design.reconcile_model_capabilities(
                    payload.get("lessonPlan") or payload,
                    payload.get("model"),
                )
                if lesson_plan.get("courseProfile", {}).get("stage") == "本科":
                    design_quality = lesson_design.validate_lesson(lesson_plan, require_approval=True)
                    if design_quality["blocking"]:
                        self.send_json({"error": "本科五步教案尚未通过审核", "quality": design_quality}, 400)
                        return
                manifest = courseware.build_manifest(
                    lesson_plan,
                    payload.get("model"),
                    payload.get("questions"),
                )
                self.send_json({"manifest": manifest, "quality": manifest["quality"]})
            elif request_path == "/api/courseware/regenerate-scene":
                manifest = payload.get("manifest") or {}
                scene_id = str(payload.get("sceneId") or "")
                if not scene_id:
                    raise ValueError("缺少场景编号")
                updated = courseware.regenerate_scene(manifest, scene_id)
                self.send_json({"manifest": updated, "quality": updated["quality"]})
            elif request_path == "/api/courseware/export-zip":
                manifest = payload.get("manifest") or {}
                quality = courseware.quality(manifest)
                if quality["blocking"]:
                    self.send_json({"error": "互动课件质检发现阻断问题", "quality": quality}, 400)
                    return
                output = courseware.create_offline_zip(manifest)
                self.send_json({"fileName": output.name, "downloadUrl": f"/api/courseware/download/{output.name}", "quality": quality})
            elif request_path == "/api/courseware/prepare-views":
                self.send_json(courseware.prepare_model_views(payload.get('manifest') or {}, bool(payload.get('checkOnly'))))
            elif request_path == "/api/courseware/export-pptx":
                manifest = payload.get("manifest") or {}
                quality = courseware.quality(manifest)
                if quality["blocking"]:
                    self.send_json({"error": "互动课件质检发现阻断问题", "quality": quality}, 400)
                    return
                output, export_result = courseware.create_ppt(manifest)
                self.send_json({"fileName": output.name, "downloadUrl": f"/api/courseware/download/{output.name}", "quality": quality, "export": export_result, "sceneCount": export_result.get("sceneCount"), "slideCount": export_result.get("slideCount"), "aligned": export_result.get("aligned", False), "modelViewCount": export_result.get("modelViewCount", 0), "warnings": export_result.get("warnings", [])})
            elif request_path == "/api/lesson/preview":
                self.send_json({"lessonPlan": build_lesson(payload)})
            elif request_path == "/api/biomed/model/generate":
                self.send_json({"error": "内置生物医学模型生成已停用，上传图片建模必须使用混元 3D"}, 410)
            elif request_path == "/api/course/generate-bundle":
                plan = payload.get("lessonPlan") or build_lesson(payload)
                if plan.get("courseProfile", {}).get("stage") != "本科":
                    raise ValueError("此接口用于本科生物医学课程")
                quality = quality_check(plan)
                if quality["blocking"]:
                    self.send_json({"error": "课件质检发现阻断问题", "quality": quality}, 400)
                    return
                output = create_ppt(plan)
                raise RuntimeError("整套生成接口不再自动调用内置模型；请在当前课程中上传图片并使用 混元 3D 建模")
            elif request_path == "/api/ppt/preview":
                plan = payload.get("lessonPlan") or build_lesson(payload)
                slides = expanded_slides(plan)
                self.send_json({"slides": [{**slide, "index": index, "previewSvg": render_preview_svg(slide, index, plan)} for index, slide in enumerate(slides, 1)], "quality": quality_check(plan), "visualProfile": plan.get("visualProfile") or visual_profile(plan.get("request") or {})})
            elif request_path == "/api/ppt/regenerate-slide":
                plan = payload.get("lessonPlan") or build_lesson(payload)
                index = int(payload.get("slideIndex", -1))
                slides = plan.get("slides") or []
                if index < 0 or index >= len(slides):
                    raise ValueError("页面序号超出范围")
                current = dict(slides[index])
                layout_cycle = ["hero-stack", "question-split", "three-column", "source-quote", "layer-map", "step-flow", "compare-duo", "task-timer", "evidence-list", "answer-rail", "exam-map", "observation-stage", "loop-summary", "take-home"]
                if not (plan.get("request") or {}).get("materialText"):
                    layout_cycle = [layout for layout in layout_cycle if layout != "source-quote"]
                requested_layout = payload.get("layout")
                if requested_layout and requested_layout not in layout_cycle:
                    raise ValueError("未知页面版式")
                if requested_layout:
                    current["layout"] = requested_layout
                elif current.get("layout") in layout_cycle:
                    start = layout_cycle.index(current.get("layout"))
                    previous_layout = slides[index - 1].get("layout") if index > 0 else None
                    next_layout = slides[index + 1].get("layout") if index + 1 < len(slides) else None
                    for offset in range(1, len(layout_cycle) + 1):
                        candidate = layout_cycle[(start + offset) % len(layout_cycle)]
                        if candidate not in (previous_layout, next_layout):
                            current["layout"] = candidate
                            break
                else:
                    current["layout"] = "content-list"
                current["visual"] = current.get("visual") or "使用更强的课堂证明对象"
                slides[index] = current
                plan["slides"] = slides
                preview_slides = expanded_slides(plan)
                self.send_json({"lessonPlan": plan, "slides": [{**slide, "index": number, "previewSvg": render_preview_svg(slide, number, plan)} for number, slide in enumerate(preview_slides, 1)], "quality": quality_check(plan)})
            elif request_path == "/api/assets/search":
                self.send_json({"query": payload.get("query", ""), "results": search_assets(payload.get("query", ""), payload.get("limit", 6))})
            elif request_path == "/api/assets/cache":
                self.send_json({"asset": cache_asset(payload.get("url", ""), payload.get("metadata"))})
            elif request_path == "/api/ppt/generate":
                plan = payload.get("lessonPlan") or build_lesson(payload)
                quality = quality_check(plan)
                if quality["blocking"]:
                    self.send_json({"error": "课件质检发现阻断问题", "quality": quality}, 400)
                    return
                output = create_ppt(plan)
                export_result = {"status": "direct", "verified": True, "message": "已生成 PPTX，浏览器已开始下载"}
                self.send_json({"fileName": output.name, "downloadUrl": f"/api/ppt/download/{output.name}", "quality": quality, "export": export_result})
            else:
                self.send_json({"error": "Not found"}, 404)
        except (ValueError, KeyError, zipfile.BadZipFile) as exc:
            print(f"[teacher-studio] request rejected: {exc}", flush=True)
            self.send_json({"error": str(exc)}, 400)
        except Exception as exc:
            print(f"[teacher-studio] request failed: {exc}", flush=True)
            self.send_json({"error": str(exc)}, 500)

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path in ("/api/lesson-ai/status", "/api/local-model/status"):
            self.send_json(lesson_ai_adapter.status())
            return
        if parsed.path == "/api/biomed/models":
            self.send_json({"models": biomed_models.catalog(), "blenderAvailable": biomed_models.find_blender() is not None})
            return
        for model_prefix in ("/api/3d/download/", "/api/3d/preview/"):
            if parsed.path.startswith(model_prefix):
                path = (OUTPUT_DIR / unquote(parsed.path[len(model_prefix):])).resolve()
                allowed = {".blend": "application/octet-stream", ".glb": "model/gltf-binary", ".png": "image/png"}
                if OUTPUT_DIR.resolve() not in path.parents or path.suffix.lower() not in allowed or not path.is_file():
                    self.send_json({"error": "资源不存在"}, 404)
                    return
                self.send_bytes(path.read_bytes(), allowed[path.suffix.lower()])
                return
        if parsed.path == "/api/3d/status":
            hunyuan_base, hunyuan_key, hunyuan_model = hunyuan3d_config()
            express_key = env_value("HUNYUAN3D_EXPRESS_API_KEY", hunyuan_key)
            segmentation_status = partfield_adapter.status()
            if hunyuan_key:
                self.send_json({
                    **local_reconstruction.status(biomed_models.find_blender() is not None),
                    "provider": "hunyuan3d",
                    "paidApi": True,
                    "networkRequired": True,
                    "apiBase": hunyuan_base,
                    "model": hunyuan_model,
                    "models": [
                        {"id": "hy-3d-3.1", "label": "混元 3D 3.1", "available": bool(hunyuan_key)},
                        {"id": "hy-3d-express", "label": "混元 3D Express 极速版", "available": bool(express_key)},
                    ],
                    "segmentation": segmentation_status,
                    "singleImage": {"available": True, "engine": "hunyuan3d", "quality": "hunyuan3d-high-precision"},
                    "multiview": {
                        "available": True,
                        "engine": "hunyuan3d-four-view",
                        "quality": "hunyuan3d-high-precision",
                        "fusionAvailable": True,
                        "scope": "base64-upload",
                        "genericFusionAvailable": True,
                        "inputViewCount": 4,
                        "providerViewCapacity": 8,
                        "reason": "混元 3D 3.1 支持多视角输入，本页面固定提交正面、左侧、右侧和背面四张",
                    },
                    "multiviewViews": list(MULTIVIEW_KEYS),
                })
                return
            self.send_json({
                **local_reconstruction.status(biomed_models.find_blender() is not None),
                "provider": "hunyuan3d",
                "configured": False,
                "models": [
                    {"id": "hy-3d-3.1", "label": "混元 3D 3.1", "available": False},
                    {"id": "hy-3d-express", "label": "混元 3D Express 极速版", "available": False},
                ],
                "segmentation": segmentation_status,
                "singleImage": {"available": False, "engine": "hunyuan3d", "quality": "unavailable"},
                "multiview": {"available": False, "engine": "hunyuan3d-four-view", "quality": "unavailable", "fusionAvailable": False, "inputViewCount": 4, "providerViewCapacity": 8},
                "multiviewViews": list(MULTIVIEW_KEYS),
            })
            return
        prefix = "/api/ppt/download/"
        if parsed.path.startswith(prefix):
            file_path = OUTPUT_DIR / unquote(parsed.path[len(prefix):])
            file_path = file_path.resolve()
            if OUTPUT_DIR.resolve() in file_path.parents and file_path.suffix == ".pptx" and file_path.is_file():
                data = file_path.read_bytes(); self.send_response(200); self._headers("application/vnd.openxmlformats-officedocument.presentationml.presentation", len(data)); self.send_header("Content-Disposition", f'attachment; filename="{file_path.name}"'); self.end_headers(); self.wfile.write(data); return
        courseware_prefix = "/api/courseware/download/"
        if parsed.path.startswith(courseware_prefix):
            file_path = (OUTPUT_DIR / unquote(parsed.path[len(courseware_prefix):])).resolve()
            allowed = {
                ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                ".zip": "application/zip",
            }
            content_type = allowed.get(file_path.suffix.lower())
            if content_type and OUTPUT_DIR.resolve() in file_path.parents and file_path.is_file():
                data = file_path.read_bytes(); self.send_response(200); self._headers(content_type, len(data)); self.send_header("Content-Disposition", f'attachment; filename="{file_path.name}"'); self.end_headers(); self.wfile.write(data); return
            self.send_json({"error": "互动课件文件不存在"}, 404)
            return
        asset_prefix = "/api/assets/"
        if parsed.path.startswith(asset_prefix):
            relative = unquote(parsed.path[len(asset_prefix):]).replace("/", os.sep)
            asset_path = (ASSET_DIR / relative).resolve()
            if asset_path.is_file() and ASSET_DIR.resolve() in asset_path.parents:
                content_type = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".gif": "image/gif"}.get(asset_path.suffix.lower(), "application/octet-stream")
                self.send_bytes(asset_path.read_bytes(), content_type)
                return
        self.send_json({"service": "teacher-studio", "status": "ok"})


if __name__ == "__main__":
    port = int(os.environ.get("TEACHER_STUDIO_PORT", "8765"))
    host = os.environ.get("TEACHER_STUDIO_HOST", "127.0.0.1")
    server = ThreadingHTTPServer((host, port), Handler)
    print(f"Teacher Studio API listening on http://{host}:{port}", flush=True)
    server.serve_forever()

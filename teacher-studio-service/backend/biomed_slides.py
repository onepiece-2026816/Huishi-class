"""One measured scene description for editable PPTX and browser SVG previews."""
import html
import re
import uuid

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

BG, INK, ACCENT, SURFACE, MUTED = "101416", "F4F7F8", "60E1C2", "232C30", "C1CDD1"
LAYOUTS = ["hero-stack", "question-split", "three-column", "source-quote", "layer-map", "step-flow", "compare-duo", "task-timer", "evidence-list", "answer-rail", "exam-map", "observation-stage", "loop-summary", "take-home"]


def wrap(value, width, size):
    lines, line, used = [], "", 0
    for char in str(value):
        advance = size * (1 if ord(char) > 255 else 0.58)
        if char == "\n" or (line and used + advance > width):
            lines.append(line)
            line, used = "", 0
        if char != "\n":
            line += char
            used += advance
    if line:
        lines.append(line)
    return lines or [""]


def scene(slide, index, plan):
    shapes = []
    def rect(x,y,w,h,fill=SURFACE,kind="rect"):
        shapes.append(dict(kind=kind,x=x,y=y,w=w,h=h,fill=fill))
    def text(value,x,y,w,h,size=26,fill=INK,bold=False):
        shapes.append(dict(kind="text",x=x,y=y,w=w,h=h,size=size,fill=fill,bold=bold,lines=wrap(value,w,size)))
    def block(value,x,y,w,h,number=None):
        rect(x,y,w,h)
        if number is not None:
            text(f"{number:02d}",x+20,y+16,w-40,36,28,ACCENT,True)
        text(value,x+20,y+(64 if number is not None else 24),w-40,h-(80 if number is not None else 40))
    rect(0,0,1280,720,BG)
    rect(0,0,8,720,ACCENT)
    text("本科 / 生物医学",56,25,500,24,17,ACCENT,True)
    text(f"{index:02d}",1160,25,64,30,22,ACCENT,True)
    text(slide.get("title",""),56,78,1168,110,38,INK,True)
    rect(56,198,74,4,ACCENT)
    values = slide.get("items",[])
    layout = slide.get("layout")
    if layout == "hero-stack":
        rect(890,236,300,300,ACCENT,"ellipse")
        text("结构\n机制\n证据",965,279,180,230,42,BG,True)
        for i,value in enumerate(values):
            text(value,60,255+i*114,760,105,30 if i == 0 else 26)
    elif layout == "question-split":
        text(values[0] if values else "",60,249,730,280,36,INK,True)
        block(values[1] if len(values)>1 else slide.get("purpose",""),850,248,370,302,1)
    elif layout in ("three-column", "step-flow"):
        for i,value in enumerate(values):
            if len(values) > 3:
                block(value,56+(i%2)*596,230+(i//2)*200,568,188,i+1)
            else:
                block(value,56+i*396,250,366,322,i+1)
            if layout == "step-flow" and len(values) <= 3 and i < len(values)-1:
                rect(424+i*396,386,24,24,ACCENT,"arrow")
    elif layout in ("compare-duo", "source-quote"):
        for i,value in enumerate(values):
            if len(values) == 3:
                block(value,56+i*396,244,366,340,i+1)
            else:
                block(value,56+i*596,244,568,340,i+1)
    elif layout in ("layer-map", "exam-map"):
        for i,value in enumerate(values):
            x = 56+i*72
            rect(x,234+i*110,1090-i*72,84,SURFACE)
            rect(x,234+i*110,8,84,ACCENT)
            text(value,x+30,254+i*110,1030-i*72,66,27)
    elif layout == "observation-stage":
        for i,value in enumerate(values):
            block(value,56 if i == 0 else 674,238 if i < 2 else 431,580 if i == 0 else 550,374 if i == 0 else 178,i+1)
    elif layout == "evidence-list":
        block(values[0] if values else "",56,238,618,370,1)
        for i,value in enumerate(values[1:]):
            block(value,704,238+i*190,520,172,i+2)
    elif layout == "task-timer":
        text("03 MIN",60,236,300,68,44,ACCENT,True)
        for i,value in enumerate(values):
            text(f"{i+1}. {value}",380,239+i*120,820,112,27)
    elif layout == "answer-rail":
        rect(81,249,4,325,ACCENT)
        for i,value in enumerate(values):
            rect(60,250+i*114,44,44,ACCENT,"ellipse")
            text(str(i+1),72,256+i*114,30,34,24,BG,True)
            text(value,138,250+i*114,1060,102,28)
    elif layout == "loop-summary":
        for i,value in enumerate(values):
            rect(56+i*396,248,350,8,ACCENT if i != 1 else "FFB469")
            text(value,56+i*396,295,350,256,30)
    else:
        for i,value in enumerate(values):
            text(f"{i+1:02d}",56,248+i*112,64,48,30,ACCENT,True)
            text(value,145,248+i*112,1078,105,28)
    source_ids = slide.get("sourceIds",[])
    refs = plan.get("references",[])
    source_names = [refs[i]["title"] for i in source_ids if 0 <= i < len(refs)]
    # Full attribution is preserved in editable speaker notes and the reference page.
    text("依据：" + "；".join(source_names),56,642,1168,48,16,MUTED)
    return shapes


def render_svg(slide,index,plan):
    parts = ['<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720">']
    for s in scene(slide,index,plan):
        x,y,w,h = (s[k] for k in ("x","y","w","h"))
        fill = s["fill"]
        if s["kind"] == "text":
            for n,line in enumerate(s["lines"]):
                parts.append(f'<text x="{x}" y="{y+s["size"]+n*s["size"]*1.3}" font-family="Microsoft YaHei,sans-serif" font-size="{s["size"]}" font-weight="{700 if s["bold"] else 400}" fill="#{fill}">{html.escape(line)}</text>')
        elif s["kind"] == "ellipse":
            parts.append(f'<ellipse cx="{x+w/2}" cy="{y+h/2}" rx="{w/2}" ry="{h/2}" fill="#{fill}"/>')
        elif s["kind"] == "arrow":
            parts.append(f'<path d="M{x} {y+h*.3}H{x+w*.55}V{y}L{x+w} {y+h/2}L{x+w*.55} {y+h}V{y+h*.7}H{x}Z" fill="#{fill}"/>')
        else:
            parts.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="#{fill}"/>')
    return "".join(parts)+"</svg>"


def quality(plan):
    blocking, warnings = [], []
    slides = plan.get("slides",[])
    if not slides:
        blocking.append(dict(code="empty-deck",message="课件为空"))
    for index,slide in enumerate(slides,1):
        if not slide.get("title") or not slide.get("items"):
            blocking.append(dict(code="empty-slide",slide=index,message="页面标题或内容为空"))
        for shape in scene(slide,index,plan):
            if shape["x"] < 0 or shape["y"] < 0 or shape["x"]+shape["w"] > 1280 or shape["y"]+shape["h"] > 720:
                blocking.append(dict(code="shape-overflow",slide=index,message="元素超出页面边界，请拆分页内容或切换版式"))
                break
            if shape["kind"] == "text" and len(shape["lines"])*shape["size"]*1.3 > shape["h"]+2:
                blocking.append(dict(code="text-overflow",slide=index,message="文本超出安全区域，请拆分此页或缩短内容"))
                break
        if index>1 and slide.get("layout") == slides[index-2].get("layout"):
            warnings.append(dict(code="repeated-layout",slide=index,message="连续页面版式相同"))
    for field,label in (("learningOutcomes","学习成果"),("prerequisites","先修知识"),("assessmentPlan","考核设计"),("references","参考来源")):
        if not plan.get(field):
            blocking.append(dict(code="missing-"+field,message="缺少"+label))
    for kind in ("知识结构","观察","案例","练习","总结"):
        if not any(s.get("type")==kind for s in slides):
            warnings.append(dict(code="missing-"+kind,message="缺少"+kind+"页面"))
    if not any(s.get("experienceMode") for s in slides):
        warnings.append(dict(code="missing-interaction",message="缺少可操作的模型观察或课堂挑战"))
    if not any(s.get("feedback") for s in slides):
        warnings.append(dict(code="missing-feedback",message="缺少操作后的即时反馈规则"))
    model = plan.get("modelPlan",{})
    if not model.get("observationTasks"):
        blocking.append(dict(code="no-observation",message="缺少结构观察任务"))
    text = " ".join(str(v) for s in slides for v in s.get("items",[]))
    for part in model.get("partDetails",[]):
        if part["name"] not in text:
            warnings.append(dict(code="missing-label-"+part["key"],message="PPT缺少模型标签："+part["name"]))
    order = [int(s.get("phaseId","phase-0").split("-")[-1]) for s in slides]
    if order != sorted(order):
        warnings.append(dict(code="phase-order",message="PPT顺序与教案流程不一致"))
    for ref in plan.get("references",[]):
        if not ref.get("title") or not ref.get("sourceType") or (ref.get("sourceType")=="web-reference" and not ref.get("sourceUrl")):
            warnings.append(dict(code="reference-source",message="参考资料缺少来源信息"))
    if plan.get("request",{}).get("materialPolicy")=="严格依据教材":
        warnings.append(dict(code="template-supplement",message="当前本科模板含补充内容，不是严格教材生成；请核对后再导出"))
    return dict(score=max(0,100-35*len(blocking)-5*len(warnings)),blocking=blocking,warnings=warnings,slideCount=len(slides),layouts=sorted({s.get("layout","") for s in slides}))


def create_ppt(plan,output_dir):
    prs = Presentation()
    prs.slide_width,prs.slide_height = Inches(13.333333),Inches(7.5)
    def unit(px):
        return Inches(px/96)
    for index,data in enumerate(plan["slides"],1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for s in scene(data,index,plan):
            args = [unit(s[k]) for k in ("x","y","w","h")]
            if s["kind"] == "text":
                tf = slide.shapes.add_textbox(*args).text_frame
                tf.word_wrap = False
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                for n,line in enumerate(s["lines"]):
                    p=tf.paragraphs[0] if n==0 else tf.add_paragraph()
                    p.text=line
                    p.line_spacing=Pt(s["size"]*0.75*1.3)
                    p.space_before=p.space_after=Pt(0)
                    for run in p.runs:
                        run.font.name="Microsoft YaHei"
                        run.font.size=Pt(s["size"]*0.75)
                        run.font.bold=s["bold"]
                        run.font.color.rgb=RGBColor.from_string(s["fill"])
            else:
                kind={"ellipse":MSO_SHAPE.OVAL,"arrow":MSO_SHAPE.RIGHT_ARROW}.get(s["kind"],MSO_SHAPE.RECTANGLE)
                shape=slide.shapes.add_shape(kind,*args)
                shape.fill.solid()
                shape.fill.fore_color.rgb=RGBColor.from_string(s["fill"])
                shape.line.fill.background()
        refs=plan.get("references",[])
        citations=[refs[i] for i in data.get("sourceIds",[]) if 0<=i<len(refs)]
        notes=["教学意图："+data.get("purpose",""),"互动任务："+data.get("interaction", "")]
        if data.get("type")=="观察":
            notes += ["模型页码："+str(index), "；".join(plan["modelPlan"]["observationTasks"]), plan["modelPlan"]["quality"]]
        notes += [f"来源：{r['title']} | {r['author']} | {r['year']} | {r['license']} | {r['sourceUrl']}" for r in citations]
        slide.notes_slide.notes_text_frame.text="\n".join(notes)
    filename="biomed_"+uuid.uuid4().hex[:10]+".pptx"
    path=output_dir/filename
    prs.save(path)
    return path

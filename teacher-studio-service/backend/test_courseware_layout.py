import json
import re
import unittest
import tempfile
from pathlib import Path
from PIL import Image
from unittest.mock import patch

from pptx import Presentation

import courseware


class CoursewareLayoutTest(unittest.TestCase):
    def test_four_views_on_activity_and_challenge_without_preview(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = {}
            for index, key in enumerate(('front', 'left', 'right', 'back')):
                paths[key] = Path(directory) / (key + '.png')
                Image.new('RGB', (60, 90), (index * 50, 80, 100)).save(paths[key])
            manifest = courseware.build_manifest({'title': '模型测试', 'activities': [{'name': '观察'}]},
                {'modelUrl': '/api/3d/download/test/model.glb', 'viewImageUrls': {key: str(path) for key, path in paths.items()}})
            with patch.object(courseware, 'prepare_model_views', return_value={'status': 'ready', 'viewImageUrls': paths}), patch.object(courseware, '_resolve_image', side_effect=lambda url: Path(url) if url else None):
                output, report = courseware.create_ppt(manifest)
            self.addCleanup(output.unlink, missing_ok=True)
            ppt = Presentation(output)
            self.assertEqual(report['modelViewCount'], 4)
            self.assertEqual(report['status'], 'direct')
            self.assertTrue(report['verified'])
            for slide, scene in zip(ppt.slides, manifest['scenes']):
                pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
                self.assertEqual(len(pictures), 4 if scene.get('model') is not None else 0)
                for picture in pictures:
                    self.assertLess(picture.left + picture.width, 7 * 914400)
            with self.assertRaises(RuntimeError):
                courseware._verify_embedded_views(output, {1: list(paths.values())})

    def test_missing_model_views_fail_before_export(self):
        manifest = {'title': '其他课程', 'model': {'viewImageUrls': {'front': '/images/missing.png'}},
                    'scenes': [{'model': {'camera': 'free'}}]}
        with self.assertRaises(ValueError):
            courseware.create_ppt(manifest)

    def test_heart_references_cannot_replace_missing_model(self):
        scene = {'model': {'camera': 'free'}}
        with self.assertRaises(ValueError):
            courseware._ppt_view_images({'title': '心脏结构', 'scenes': [scene]})
        with self.assertRaises(ValueError):
            courseware._ppt_view_images({'title': '心脏结构', 'model': {'viewImageUrls': {'front': '/images/missing.png'}}, 'scenes': [scene]})
        self.assertEqual(courseware._ppt_view_images({'title': '其他课程', 'scenes': []}), {})

    def test_manifest_preserves_model_view_image_urls(self):
        manifest = courseware.build_manifest({
            'id': 'views-test',
            'title': '视角测试',
            'request': {'topic': '结构观察'},
            'modelAsset': {
                'modelUrl': '/api/3d/download/test/model.glb',
                'previewUrl': '/api/3d/preview/test/preview.png',
                'viewImageUrls': {
                    'front': '/api/3d/preview/test/front.png',
                    'left': '/api/3d/preview/test/left.png',
                    'right': '/api/3d/preview/test/right.png',
                    'back': '/api/3d/preview/test/back.png',
                },
            },
        })
        self.assertEqual(manifest['model']['viewImageUrls']['front'], '/api/3d/preview/test/front.png')
        self.assertEqual(manifest['model']['viewImageUrls']['back'], '/api/3d/preview/test/back.png')

    def test_web_and_ppt_share_scene_content_contract(self):
        scene = {
            'type': 'question-chain',
            'claim': '核心结论',
            'teacherCue': '教师追问',
            'studentAction': '学生任务',
            'expectedOutput': '预期产出',
            'detail': {
                'learningContent': '学习内容',
                'studentSteps': ['步骤一', '步骤二'],
                'questioning': ['追问一'],
                'evaluationRubric': ['评价标准'],
                'remediation': '补救措施',
            },
        }
        blocks = courseware.scene_content_blocks(scene)
        self.assertEqual([item['label'] for item in blocks], ['学习内容', '学习步骤', '教师追问', '评价标准'])
        self.assertEqual([item['value'] for item in blocks], ['学习内容', '步骤一；步骤二', '追问一', '评价标准'])

    def test_generated_copy_is_chinese_and_long_ppt_is_paginated(self):
        plan = {
            'id': 'layout-test', 'title': '心脏结构与功能',
            'objectives': ['依据结构图解释血流路径'],
            'overview': '心房接收血液，心室推动血液进入循环。' * 80,
            'activities': [{'name': '结构定位', 'learningContent': '比较心房与心室的连接关系。' * 80,
                            'studentSteps': ['标记结构', '绘制路径', '解释依据'],
                            'teacherProcess': '先观察，再追问依据。', 'visibleOutput': '带依据的路径图',
                            'evaluation': {'standard': '关键结构和方向正确'}}],
            'homework': '提交结构图和机制解释。',
        }
        manifest = courseware.build_manifest(plan)
        for scene in manifest['scenes']:
            for field in ('title', 'claim', 'teacherCue', 'studentAction', 'expectedOutput'):
                self.assertFalse(re.search(r'[A-Za-z]{3,}', scene[field]), (scene['id'], field))
        self.assertFalse(any('heart-anatomy' in asset['url'] for asset in manifest['assets']))
        output, report = courseware.create_ppt(manifest)
        self.addCleanup(output.unlink, missing_ok=True)
        ppt = Presentation(output)
        self.assertEqual(len(ppt.slides), len([scene for scene in manifest['scenes'] if scene.get('enabled', True)]))
        self.assertEqual(report['status'], 'direct')
        self.assertEqual(report['slideCount'], len(ppt.slides))
        texts = []
        for slide in ppt.slides:
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.left + shape.width, ppt.slide_width + 1)
                self.assertLessEqual(shape.top + shape.height, ppt.slide_height + 1)
                if not shape.has_text_frame:
                    continue
                texts.append(shape.text)
                self.assertFalse(re.search(r'[A-Za-z]{3,}', shape.text), shape.text)
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        self.assertEqual(run.font.name, '宋体')
                        self.assertIn('typeface="宋体"', run._r.xml)
        self.assertIn('比较心房与心室的连接关系。', '\n'.join(texts))

    def test_saved_english_scene_is_localized_for_export(self):
        scene = {'title': 'Observe the model and explain the structure',
                 'detail': {'studentSteps': ['Complete and submit the transfer task.']}}
        translated = courseware._chinese_copy(scene)
        self.assertEqual(translated['title'], '观察结构并解释功能')
        self.assertEqual(translated['detail']['studentSteps'], ['完成并提交迁移任务。'])
        self.assertEqual(scene['title'], 'Observe the model and explain the structure')


if __name__ == '__main__':
    unittest.main()
